# /// script
# requires-python = ">=3.11"
# dependencies = [
#   "m365-gdap-runner @ git+https://github.com/DrollRobot/m365-gdap-runner.git@v0.1.0",
#   "openpyxl>=3.1.5",
# ]
# ///
"""Example script: the Entra sign-in logs that match a filter, in each tenant.

Reads ``/auditLogs/signIns`` with the OData ``$filter`` given. ``--beta`` reads the beta
endpoint::

    uv run entra_sign_in_logs.py "userPrincipalName eq 'admin@contoso.com'" --include pilot
    uv run entra_sign_in_logs.py "signInEventTypes/any(t: t eq 'nonInteractiveUser')" --beta

Writes one JSON file per tenant with matches, and one workbook of match counts per tenant
with the filter in its title row, in the current folder or ``--out``. A file never replaces
another: where its name is taken, it gets the first free ``_2``, ``_3``, and so on.

Needs ``AuditLog.Read.All``, a GDAP role such as Reports Reader, and Entra ID P1 or P2.

Common filters. Text in single quotes, dates and numbers bare, joined with ``and``/``or``::

    # Sign-in type (beta). Without one, only interactive user sign-ins return.
    signInEventTypes/any(t: t eq 'interactiveUser')
    signInEventTypes/any(t: t eq 'nonInteractiveUser')
    signInEventTypes/any(t: t eq 'servicePrincipal')
    signInEventTypes/any(t: t eq 'managedIdentity')
    signInEventTypes/any(t: t eq 'interactiveUser' or t eq 'nonInteractiveUser')

    # Date range: ge, le or eq, UTC. Entra keeps 30 days.
    createdDateTime ge 2026-10-01T00:00:00Z
    createdDateTime ge 2026-10-01T00:00:00Z and createdDateTime le 2026-10-07T23:59:59Z

    # Success, wrong password, any failure (ne is not in Microsoft's docs).
    status/errorCode eq 0
    status/errorCode eq 50126
    status/errorCode ne 0

    # IP address: eq or startsWith.
    ipAddress eq '203.0.113.10'
    startsWith(ipAddress, '203.0.113.')

    # Autonomous system number (beta; not in Microsoft's docs as filterable).
    autonomousSystemNumber eq 64500

    # User agent (beta): eq or startsWith.
    startsWith(userAgent, 'python-requests')

Microsoft's docs: listing sign-ins, sign-in properties and their operators (beta, v1.0),
``$filter`` syntax, error codes, retention::

    https://learn.microsoft.com/en-us/graph/api/signin-list?view=graph-rest-beta
    https://learn.microsoft.com/en-us/graph/api/resources/signin?view=graph-rest-beta
    https://learn.microsoft.com/en-us/graph/api/resources/signin?view=graph-rest-1.0
    https://learn.microsoft.com/en-us/graph/filter-query-parameter
    https://learn.microsoft.com/en-us/entra/identity-platform/reference-error-codes
    https://login.microsoftonline.com/error
    https://learn.microsoft.com/en-us/entra/identity/monitoring-health/reference-reports-data-retention
"""

from __future__ import annotations

import datetime
import json
import logging
import re
from pathlib import Path
from typing import BinaryIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from m365_gdap_runner import TenantContext, TenantResult

logger = logging.getLogger(__name__)

CLIENTS = ["graph"]

# The one scope Graph asks of a delegated sign-in that lists sign-ins.
SCOPES = ["AuditLog.Read.All"]

# A report only reads, so one tenant's error is no reason to stop the others.
CONTINUE_ON_ERROR = True

# The OData filter sent as ``$filter``. The command line sets it before the run.
FILTER = ""

# Read Graph's beta endpoint instead of v1.0. ``--beta`` sets it before the run.
BETA = False

# Where the JSON files and the workbook go: the current folder. ``--out`` replaces it.
OUT_DIR = Path()

# Output files are named this, then the tenant for a JSON file, then today's date.
FILE_STEM = "EntraSignInLogs"

# The workbook's columns, in order, below the title row.
COLUMNS = ["Tenant", "TenantId", "Status", "SignIns", "Error"]

# Widest a column is made to fit its longest value, in characters.
MAX_COLUMN_WIDTH = 80


def run(ctx: TenantContext) -> int:
    """Read the tenant's matching sign-ins, and write them as JSON when there are any.

    Args:
        ctx: The tenant context supplied by the runner.

    Returns:
        How many sign-ins match.
    """
    sign_ins = ctx.graph.get_all(_collection(ctx), params={"$filter": FILTER})
    if not sign_ins:
        logger.info("No sign-ins match")
        return 0
    path, file = _create(f"{FILE_STEM}_{_file_safe(ctx.tenant.display_name)}_{_today()}", ".json")
    with file:
        file.write(json.dumps(sign_ins, indent=2).encode("utf-8"))
    logger.info("%d sign-ins match; wrote %s", len(sign_ins), path)
    return len(sign_ins)


def report(results: list[TenantResult]) -> None:
    """Write one workbook with each tenant's count of matches, under the filter as its title.

    Args:
        results: Every tenant's result, in run order. A tenant that failed has its status
            and error, and a blank count.
    """
    if not results:
        print("No tenants ran; no workbook written.")
        return
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None  # noqa: S101 (a new workbook always has a sheet)
    sheet.title = "Tenants"
    sheet.append([FILTER])
    sheet["A1"].font = Font(bold=True)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    sheet.append(COLUMNS)
    for result in results:
        sheet.append(
            [
                result.tenant.display_name,
                result.tenant.tenant_id,
                result.status,
                result.result if isinstance(result.result, int) else None,
                result.error,
            ]
        )
    last_cell = sheet.cell(sheet.max_row, len(COLUMNS)).coordinate
    table = Table(displayName="Tenants", ref=f"A2:{last_cell}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium8", showRowStripes=True)
    sheet.add_table(table)
    sheet.freeze_panes = "A3"
    # The title row spans the table, so it is left out of the widths.
    for number, column in enumerate(sheet.iter_cols(min_row=2, values_only=True), start=1):
        longest = max(len(str(value)) for value in column if value is not None)
        width = min(longest + 2, MAX_COLUMN_WIDTH)
        sheet.column_dimensions[get_column_letter(number)].width = width

    path, file = _create(f"{FILE_STEM}_{_today()}", ".xlsx")
    with file:
        workbook.save(file)
    print(f"Wrote {path}")


def _collection(ctx: TenantContext) -> str:
    """Return the sign-ins collection to read: v1.0 by its path, beta by its whole URL.

    ``ctx.graph`` sends a path to v1.0, and an absolute URL on its own host as it is.

    Args:
        ctx: The tenant context, for the Graph address of the tenant's cloud.

    Returns:
        ``/auditLogs/signIns``, or with ``--beta`` the beta collection's URL.
    """
    if BETA:
        return f"{ctx.auth.tenant.cloud.graph}/beta/auditLogs/signIns"
    return "/auditLogs/signIns"


def _create(stem: str, suffix: str) -> tuple[Path, BinaryIO]:
    """Create a new file in :data:`OUT_DIR`, never one that is already there.

    Each name is created only if it does not exist, so neither an earlier run's file nor one
    another tenant of the same name just wrote is replaced.

    Args:
        stem: The file name before its suffix, such as ``EntraSignInLogs_2026-10-29``.
        suffix: The file name's suffix, such as ``.xlsx``.

    Returns:
        Where the file is, and the file, open for writing bytes. The name is *stem* and
        *suffix*, or where that is taken, the first free one with ``_2``, ``_3``, and so on
        between them.
    """
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"{stem}{suffix}"
    number = 1
    while True:
        try:
            return path, path.open("xb")
        except FileExistsError:
            number += 1
            path = OUT_DIR / f"{stem}_{number}{suffix}"


def _file_safe(name: str) -> str:
    """Turn a tenant's display name into the part of a file name that names it.

    Args:
        name: The display name, such as ``Contoso Ltd.``.

    Returns:
        The name with anything but letters, digits, ``_`` and ``-`` replaced by ``_``.
    """
    return re.sub(r"[^\w-]", "_", name)


def _today() -> str:
    """Return today's date as it goes in a file name.

    Returns:
        The date, such as ``2026-10-29``.
    """
    return datetime.date.today().isoformat()


if __name__ == "__main__":
    import argparse

    from m365_gdap_runner import add_arguments, main

    parser = argparse.ArgumentParser(
        description="Write the Entra sign-in logs that match a filter, for each tenant."
    )
    parser.add_argument(
        "filter",
        metavar="FILTER",
        help="OData $filter for /auditLogs/signIns, such as "
        "\"userPrincipalName eq 'admin@contoso.com'\".",
    )
    parser.add_argument(
        "--beta",
        action="store_true",
        help="Read Graph's beta endpoint instead of v1.0.",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=OUT_DIR,
        metavar="DIR",
        help="Folder for the JSON files and the workbook (default: the current folder).",
    )
    add_arguments(parser)
    args = parser.parse_args()
    FILTER = args.filter
    BETA = args.beta
    OUT_DIR = args.out
    raise SystemExit(main(__name__, args))
