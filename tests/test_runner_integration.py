"""Integration tests: runner with the real session, azure-auth clients and the examples."""

from __future__ import annotations

import argparse
import datetime
import json
import logging
import os
from pathlib import Path
from typing import Any

import httpx2
import pytest
from azure_auth.constants import EXCHANGE_POWERSHELL_CLIENT_ID, GRAPH_CLI_CLIENT_ID
from openpyxl import load_workbook

from m365_gdap_runner import cli
from m365_gdap_runner.config.paths import CONFIG_DIR_ENV
from m365_gdap_runner.config.schema import ENV_PREFIX
from m365_gdap_runner.loader import Payload, payload_from_module
from m365_gdap_runner.runner import TenantResult, run_payload
from m365_gdap_runner.session import PartnerSession
from m365_gdap_runner.tenants import Tenant
from tests._gdap_fakes import (
    CONSENT_REQUIRED,
    PARTNER_TENANT,
    TENANT_A,
    TENANT_B,
    USER,
    FakeMsal,
    bearer_tenant,
    fake_services,
    load_script,
    make_settings,
)

pytestmark = pytest.mark.integration

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"
AUDIT_DISABLED = EXAMPLES / "audit_disabled.py"
UNIFIED_AUDIT_LOG = EXAMPLES / "unified_audit_log.py"
SERVICE_PRINCIPAL_ENABLED = EXAMPLES / "service_principal_enabled.py"
ENTRA_SIGN_IN_LOGS = EXAMPLES / "entra_sign_in_logs.py"
ENTRA_AUDIT_LOGS = EXAMPLES / "entra_audit_logs.py"
MESSAGE_TRACE = EXAMPLES / "exchange_message_trace.py"
GRAPH_SCOPES = EXAMPLES / "graph_scopes.py"


def _example(path: Path) -> Payload:
    return payload_from_module(load_script(path), path.stem)


@pytest.fixture(autouse=True)
def _isolated_config(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Strip this package's env vars and point the config dir at a temp dir."""
    for key in list(os.environ):
        if key.startswith(ENV_PREFIX):
            monkeypatch.delenv(key)
    monkeypatch.setenv(CONFIG_DIR_ENV, str(tmp_path))


def test_graph_example_runs_through_session_and_azure_auth(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    fake_msal.tenant_errors[TENANT_B] = CONSENT_REQUIRED
    reads: list[httpx2.Request] = []
    transport = _directory_audit_services({TENANT_A: [_directory_audit(1)]}, reads)

    module = load_script(ENTRA_AUDIT_LOGS)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path)
    payload = payload_from_module(module, ENTRA_AUDIT_LOGS.stem)
    session = PartnerSession(make_settings(), transport=transport)
    session.login(payload.clients)
    tenants = session.list_tenants()
    ready = session.check_access(tenants, scopes=payload.scopes, clients=payload.clients)
    results = run_payload(session, payload, tenants, ready=ready)

    assert [t.tenant_id for t in tenants] == [TENANT_A, TENANT_B]
    assert [r.status for r in results] == ["ok", "auth_failed"]
    assert results[0].result == 1
    # Consent was checked in the tenant, which has none, and beta corp has no account to try.
    assert results[1].error == "partner: no Graph consent; direct: no m365_username"
    # The script never ran for the tenant that refused the token.
    assert [bearer_tenant(r) for r in reads] == [TENANT_A]
    assert fake_msal.interactive_calls == 1


def test_exchange_example_signs_in_to_exchange_and_runs_the_cmdlet(
    fake_msal: FakeMsal, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.INFO)
    payload = _example(AUDIT_DISABLED)
    assert (payload.clients, payload.scopes, payload.report) == (("exchange",), (), None)

    session = PartnerSession(make_settings(), transport=httpx2.MockTransport(fake_services))
    session.login(payload.clients)
    tenants = session.list_tenants()
    ready = session.check_access(tenants, clients=payload.clients)
    results = run_payload(session, payload, tenants, ready=ready)

    assert [r.status for r in results] == ["ok", "ok"]
    assert [r.getMessage() for r in caplog.records if r.filename == "audit_disabled.py"] == [
        "Mailbox auditing is on",
        "Mailbox auditing is off tenant-wide (AuditDisabled is True)",
    ]
    # Graph for the tenant list, Exchange for the cmdlet: two home sign-ins.
    assert [c.client_id for c in fake_msal.calls if c.method == "interactive"] == [
        GRAPH_CLI_CLIENT_ID,
        EXCHANGE_POWERSHELL_CLIENT_ID,
    ]
    # Each customer was checked with the Exchange client id before the cmdlet ran.
    assert fake_msal.silent_calls(TENANT_A)[0] == (
        EXCHANGE_POWERSHELL_CLIENT_ID,
        ("https://outlook.office365.com/.default",),
    )


def _sheets(path: Path) -> dict[str, list[tuple[Any, ...]]]:
    """Read every sheet of a workbook as rows of values, the header row first."""
    workbook = load_workbook(path)
    return {sheet.title: list(sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets}


def _assert_consolas(path: Path) -> None:
    """Check a workbook is in Consolas throughout: every cell, and its default font."""
    for sheet in load_workbook(path).worksheets:
        cells = [cell for row in sheet.iter_rows() for cell in row]
        assert {cell.font.name for cell in cells} == {"Consolas"}
        # A theme scheme would make Excel show the theme's font, Calibri, in Consolas's place.
        assert {cell.font.scheme for cell in cells} == {None}
        # A cell nothing was written to has the workbook's default font.
        assert sheet["Z99"].font.name == "Consolas"


# --- unified_audit_log.py --------------------------------------------------------------


def _ual_services(request: httpx2.Request) -> httpx2.Response:
    """Fake Exchange for the unified audit log example: on in tenant A, off in tenant B."""
    tenant = bearer_tenant(request)
    if request.url.path != f"/adminapi/beta/{tenant}/InvokeCommand":
        return fake_services(request)
    cmdlet = json.loads(request.content)["CmdletInput"]["CmdletName"]
    assert cmdlet == "Get-AdminAuditLogConfig"
    config = {"UnifiedAuditLogIngestionEnabled": tenant == TENANT_A}
    return httpx2.Response(200, json={"value": [config]})


@pytest.fixture
def unified_audit_log(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """The example, writing its workbook to a temp dir, as ``--out`` makes it."""
    module = load_script(UNIFIED_AUDIT_LOG)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path / "reports")
    return module


def test_unified_audit_log_example_writes_one_workbook_for_every_tenant(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    unified_audit_log: Any,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    monkeypatch.setenv(ENV_PREFIX + "PARTNER_TENANT_ID", PARTNER_TENANT)
    monkeypatch.setenv(ENV_PREFIX + "USERNAME", USER)
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)

    parser = argparse.ArgumentParser()
    cli.add_arguments(parser)
    args = parser.parse_args(["--include", "all"])
    assert cli.main(unified_audit_log, args, transport=httpx2.MockTransport(_ual_services)) == 0

    (path,) = (tmp_path / "reports").iterdir()
    assert path.name == f"UnifiedAuditLog_{datetime.date.today().isoformat()}.xlsx"
    assert capsys.readouterr().out == f"Wrote {path}\n"
    assert _sheets(path) == {
        "Tenants": [
            ("Tenant", "TenantId", "Status", "UnifiedAuditLogIngestionEnabled", "Error"),
            ("Alpha", TENANT_A, "ok", True, None),
            ("beta corp", TENANT_B, "ok", False, None),
        ]
    }
    assert [
        (r.levelname, r.getMessage())
        for r in caplog.records
        if r.filename == UNIFIED_AUDIT_LOG.name
    ] == [
        ("INFO", "Unified audit log is on"),
        ("WARNING", "Unified audit log is off (UnifiedAuditLogIngestionEnabled is False)"),
    ]


def test_unified_audit_log_workbook_lists_the_tenants_that_failed(
    tmp_path: Path, unified_audit_log: Any
) -> None:
    unified_audit_log.report(
        [
            TenantResult(Tenant(TENANT_A, "Alpha"), "ok", False),
            TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
        ]
    )
    (path,) = (tmp_path / "reports").iterdir()
    sheet = load_workbook(path)["Tenants"]
    assert list(sheet.iter_rows(min_row=2, values_only=True)) == [
        ("Alpha", TENANT_A, "ok", False, None),
        ("beta corp", TENANT_B, "auth_failed", None, "AADSTS65001"),
    ]
    assert sheet.freeze_panes == "A2"
    (table,) = sheet.tables.values()
    assert table.ref == "A1:E3"
    assert table.tableStyleInfo is not None
    assert table.tableStyleInfo.name == "TableStyleMedium8"


def test_unified_audit_log_writes_no_workbook_when_no_tenant_ran(
    tmp_path: Path, unified_audit_log: Any, capsys: pytest.CaptureFixture[str]
) -> None:
    unified_audit_log.report([])
    assert not (tmp_path / "reports").exists()
    assert capsys.readouterr().out == "No tenants ran; no workbook written.\n"


def test_unified_audit_log_keeps_an_earlier_workbook(
    tmp_path: Path, unified_audit_log: Any, capsys: pytest.CaptureFixture[str]
) -> None:
    results = [TenantResult(Tenant(TENANT_A, "Alpha"), "ok", True)]
    unified_audit_log.report(results)
    (first,) = (tmp_path / "reports").iterdir()
    earlier = first.read_bytes()

    unified_audit_log.report(results)

    second = first.with_name(f"{first.stem}_2.xlsx")
    assert sorted(p.name for p in (tmp_path / "reports").iterdir()) == [first.name, second.name]
    assert first.read_bytes() == earlier
    assert capsys.readouterr().out == f"Wrote {first}\nWrote {second}\n"
    _assert_consolas(second)


# --- service_principal_enabled.py --------------------------------------------------------

BACKUP_APP_ID = "cw000000-0000-0000-0000-000000000001"
REPORTER_APP_ID = "rp000000-0000-0000-0000-000000000002"
MAIL_APP_ID = "ml000000-0000-0000-0000-000000000003"


def _sp(object_id: str, name: str, app_id: str, *, enabled: bool) -> dict[str, Any]:
    """Build a service principal as the list read returns it."""
    return {"id": object_id, "appId": app_id, "displayName": name, "accountEnabled": enabled}


def _sp_services(
    apps: dict[str, list[dict[str, Any]]], changes: list[httpx2.Request]
) -> httpx2.MockTransport:
    """Fake Graph holding *apps* per tenant, recording every PATCH and DELETE it is sent.

    A read of one service principal returns it whole: as listed, plus ``tags``, which the list
    read leaves out. A PATCH changes nothing, so a later read still shows the old state. A
    DELETE removes the app, so a later read of it fails.
    """

    def handler(request: httpx2.Request) -> httpx2.Response:
        held = apps.get(bearer_tenant(request), [])
        path = request.url.path
        if path == "/v1.0/servicePrincipals":
            assert request.url.params["$select"] == "id,appId,displayName,accountEnabled"
            assert request.url.params["$top"] == "999"
            return httpx2.Response(200, json={"value": held})
        if path.startswith("/v1.0/servicePrincipals/"):
            if request.method == "PATCH":
                changes.append(request)
                return httpx2.Response(204)
            (app,) = [app for app in held if app["id"] == path.rsplit("/", 1)[1]]
            if request.method == "DELETE":
                changes.append(request)
                held.remove(app)
                return httpx2.Response(204)
            return httpx2.Response(200, json={**app, "tags": ["WindowsAzureActiveDirectoryApp"]})
        return fake_services(request)

    return httpx2.MockTransport(handler)


@pytest.fixture
def sp_enabled(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """The example, run through ``main()`` with ``--include all``, writing to a temp dir.

    Its patterns are ``backup`` and the reporter app's ID, as the command line would set them.
    """
    monkeypatch.setenv(ENV_PREFIX + "PARTNER_TENANT_ID", PARTNER_TENANT)
    monkeypatch.setenv(ENV_PREFIX + "USERNAME", USER)
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)
    module = load_script(SERVICE_PRINCIPAL_ENABLED)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path / "reports")
    patterns = [module.pattern("backup"), module.pattern(REPORTER_APP_ID)]
    monkeypatch.setattr(module, "PATTERNS", patterns)
    return module


def _disable(monkeypatch: pytest.MonkeyPatch, module: Any) -> None:
    """Set what ``--disable`` sets."""
    monkeypatch.setattr(module, "DISABLE", True)
    monkeypatch.setattr(module, "SCOPES", ["Application.ReadWrite.All"])


def _delete(monkeypatch: pytest.MonkeyPatch, module: Any) -> None:
    """Set what ``--delete`` sets."""
    monkeypatch.setattr(module, "DELETE", True)
    monkeypatch.setattr(module, "SCOPES", ["Application.ReadWrite.All"])


def _run_all(module: Any, transport: httpx2.MockTransport) -> int:
    parser = argparse.ArgumentParser()
    cli.add_arguments(parser)
    return cli.main(module, parser.parse_args(["--include", "all"]), transport=transport)


def _output(tmp_path: Path, suffix: str = "xlsx") -> Path:
    """Return where the example writes today's workbook, or its JSON with ``json``."""
    today = datetime.date.today().isoformat()
    return tmp_path / "reports" / f"ServicePrincipalEnabled_{today}.{suffix}"


def test_sp_enabled_reports_every_match_and_changes_nothing(
    fake_msal: FakeMsal,
    sp_enabled: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    changes: list[httpx2.Request] = []
    apps = {
        TENANT_A: [
            _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True),
            _sp("sp-a2", "Reporter", REPORTER_APP_ID, enabled=False),
            _sp("sp-a3", "Mail", MAIL_APP_ID, enabled=True),
        ],
        TENANT_B: [_sp("sp-b1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=False)],
    }

    assert _run_all(sp_enabled, _sp_services(apps, changes)) == 0

    assert changes == []
    path = _output(tmp_path)
    assert [p.name for p in path.parent.iterdir()] == [path.name]
    assert capsys.readouterr().out == f"Wrote {path}\n"
    assert _sheets(path) == {
        "Tenants": [
            ("Tenant", "Status", "TenantId", "Contoso Cloud Backup", "Reporter"),
            ("Alpha", "ok", TENANT_A, "enabled", "disabled"),
            ("beta corp", "ok", TENANT_B, "disabled", None),
        ]
    }
    # Nothing is written while tenants run.
    assert [r for r in caplog.records if r.filename == SERVICE_PRINCIPAL_ENABLED.name] == []
    # A report asks only to read.
    assert (
        GRAPH_CLI_CLIENT_ID,
        ("https://graph.microsoft.com/Application.Read.All",),
    ) in fake_msal.silent_calls(TENANT_A)


def test_sp_enabled_disable_turns_sign_in_off_where_it_is_on(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, sp_enabled: Any, tmp_path: Path
) -> None:
    _disable(monkeypatch, sp_enabled)
    patches: list[httpx2.Request] = []
    apps = {
        TENANT_A: [
            _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True),
            _sp("sp-a2", "Reporter", REPORTER_APP_ID, enabled=False),
        ],
        TENANT_B: [_sp("sp-b1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True)],
    }

    assert _run_all(sp_enabled, _sp_services(apps, patches)) == 0

    assert [(bearer_tenant(p), p.url.path) for p in patches] == [
        (TENANT_A, "/v1.0/servicePrincipals/sp-a1"),
        (TENANT_B, "/v1.0/servicePrincipals/sp-b1"),
    ]
    assert all(json.loads(p.content) == {"accountEnabled": False} for p in patches)
    assert _sheets(_output(tmp_path))["Tenants"][1:] == [
        ("Alpha", "ok", TENANT_A, "disabled", "disabled"),
        ("beta corp", "ok", TENANT_B, "disabled", None),
    ]
    assert (
        GRAPH_CLI_CLIENT_ID,
        ("https://graph.microsoft.com/Application.ReadWrite.All",),
    ) in fake_msal.silent_calls(TENANT_A)


def test_sp_enabled_exports_each_match_whole_keyed_by_tenant_id(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    sp_enabled: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr(sp_enabled, "EXPORT_JSON", True)
    # With --disable, the export shows each app as the run left it.
    _disable(monkeypatch, sp_enabled)
    apps = {
        TENANT_A: [
            _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True),
            _sp("sp-a3", "Mail", MAIL_APP_ID, enabled=True),
        ]
    }

    assert _run_all(sp_enabled, _sp_services(apps, [])) == 0

    path = _output(tmp_path, "json")
    assert capsys.readouterr().out == f"Wrote {_output(tmp_path)}\nWrote {path}\n"
    backup = _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=False)
    assert json.loads(path.read_text(encoding="utf-8")) == {
        TENANT_A: [{**backup, "tags": ["WindowsAzureActiveDirectoryApp"]}],
        TENANT_B: [],
    }


def test_sp_enabled_delete_deletes_every_match(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, sp_enabled: Any, tmp_path: Path
) -> None:
    _delete(monkeypatch, sp_enabled)
    changes: list[httpx2.Request] = []
    apps = {
        TENANT_A: [
            _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True),
            _sp("sp-a2", "Reporter", REPORTER_APP_ID, enabled=False),
            _sp("sp-a3", "Mail", MAIL_APP_ID, enabled=True),
        ],
        TENANT_B: [_sp("sp-b1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=False)],
    }

    assert _run_all(sp_enabled, _sp_services(apps, changes)) == 0

    # Enabled or not, every match is deleted; the app that does not match is left alone.
    assert [(bearer_tenant(r), r.method, r.url.path) for r in changes] == [
        (TENANT_A, "DELETE", "/v1.0/servicePrincipals/sp-a1"),
        (TENANT_A, "DELETE", "/v1.0/servicePrincipals/sp-a2"),
        (TENANT_B, "DELETE", "/v1.0/servicePrincipals/sp-b1"),
    ]
    assert [app["id"] for app in apps[TENANT_A]] == ["sp-a3"]
    assert _sheets(_output(tmp_path))["Tenants"][1:] == [
        ("Alpha", "ok", TENANT_A, "deleted", "deleted"),
        ("beta corp", "ok", TENANT_B, "deleted", None),
    ]
    assert (
        GRAPH_CLI_CLIENT_ID,
        ("https://graph.microsoft.com/Application.ReadWrite.All",),
    ) in fake_msal.silent_calls(TENANT_A)


def test_sp_enabled_delete_exports_each_match_as_it_was_before(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, sp_enabled: Any, tmp_path: Path
) -> None:
    monkeypatch.setattr(sp_enabled, "EXPORT_JSON", True)
    _delete(monkeypatch, sp_enabled)
    changes: list[httpx2.Request] = []
    apps = {TENANT_A: [_sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True)]}

    # The fake fails a read of a deleted app, so this passes only if each is read first.
    assert _run_all(sp_enabled, _sp_services(apps, changes)) == 0

    assert [r.url.path for r in changes] == ["/v1.0/servicePrincipals/sp-a1"]
    backup = _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True)
    assert json.loads(_output(tmp_path, "json").read_text(encoding="utf-8")) == {
        TENANT_A: [{**backup, "tags": ["WindowsAzureActiveDirectoryApp"]}],
        TENANT_B: [],
    }


def test_sp_enabled_workbook_lists_failed_tenants_and_apps_that_share_a_name(
    sp_enabled: Any, tmp_path: Path
) -> None:
    sp_enabled.report(
        [
            TenantResult(
                Tenant(TENANT_A, "Alpha"),
                "ok",
                [
                    _sp("sp-1", "Backup", "app-1", enabled=True),
                    _sp("sp-2", "Backup", "app-2", enabled=False),
                    {"id": "sp-3", "appId": "app-3", "accountEnabled": True},
                ],
            ),
            TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
        ]
    )
    sheet = load_workbook(_output(tmp_path))["Tenants"]
    assert list(sheet.iter_rows(values_only=True)) == [
        # An app without a display name is shown by its app ID.
        ("Tenant", "Status", "TenantId", "app-3", "Backup"),
        ("Alpha", "ok", TENANT_A, "enabled", "disabled, enabled"),
        ("beta corp", "auth_failed", TENANT_B, None, None),
    ]
    assert sheet.freeze_panes == "A2"
    (table,) = sheet.tables.values()
    assert table.ref == "A1:E3"


def test_sp_enabled_writes_nothing_when_no_tenant_ran(
    sp_enabled: Any, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    sp_enabled.report([])
    assert not (tmp_path / "reports").exists()
    assert capsys.readouterr().out == "No tenants ran; nothing written.\n"


def test_sp_enabled_keeps_an_earlier_runs_files(
    monkeypatch: pytest.MonkeyPatch,
    sp_enabled: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr(sp_enabled, "EXPORT_JSON", True)
    earlier = [_output(tmp_path), _output(tmp_path, "json")]
    earlier[0].parent.mkdir(parents=True)
    for path in earlier:
        path.write_bytes(path.name.encode())
    app = _sp("sp-1", "Backup", BACKUP_APP_ID, enabled=True)

    sp_enabled.report([TenantResult(Tenant(TENANT_A, "Alpha"), "ok", [app])])

    for path in earlier:
        assert path.read_bytes() == path.name.encode()
    workbook, exported = (path.with_name(f"{path.stem}_2{path.suffix}") for path in earlier)
    assert capsys.readouterr().out == f"Wrote {workbook}\nWrote {exported}\n"
    assert json.loads(exported.read_text(encoding="utf-8")) == {TENANT_A: [app]}
    assert _sheets(workbook)["Tenants"][1] == ("Alpha", "ok", TENANT_A, "enabled")
    _assert_consolas(workbook)


def _refused_in_alpha(transport: httpx2.MockTransport) -> httpx2.MockTransport:
    """Wrap a fake Graph so Alpha refuses the app read, as without a role that may read it."""

    def handler(request: httpx2.Request) -> httpx2.Response:
        if request.url.path == "/v1.0/servicePrincipals" and bearer_tenant(request) == TENANT_A:
            error = {"code": "Authorization_RequestDenied", "message": "Insufficient privileges."}
            return httpx2.Response(403, json={"error": error})
        return transport.handle_request(request)

    return httpx2.MockTransport(handler)


def test_sp_enabled_report_goes_on_after_a_tenant_error(
    fake_msal: FakeMsal, sp_enabled: Any, tmp_path: Path
) -> None:
    apps = {TENANT_B: [_sp("sp-b1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True)]}

    assert _run_all(sp_enabled, _refused_in_alpha(_sp_services(apps, []))) == 1

    alpha, beta = _sheets(_output(tmp_path))["Tenants"][1:]
    assert alpha == ("Alpha", "error", TENANT_A, None)
    assert beta == ("beta corp", "ok", TENANT_B, "enabled")


def test_sp_enabled_disable_goes_on_after_a_tenant_error(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, sp_enabled: Any, tmp_path: Path
) -> None:
    _disable(monkeypatch, sp_enabled)
    patches: list[httpx2.Request] = []
    apps = {TENANT_B: [_sp("sp-b1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True)]}

    assert _run_all(sp_enabled, _refused_in_alpha(_sp_services(apps, patches))) == 1

    rows = _sheets(_output(tmp_path))["Tenants"][1:]
    assert [(row[0], row[1]) for row in rows] == [("Alpha", "error"), ("beta corp", "ok")]
    assert rows[1][3] == "disabled"
    assert [(bearer_tenant(p), p.url.path) for p in patches] == [
        (TENANT_B, "/v1.0/servicePrincipals/sp-b1")
    ]


def test_sp_enabled_delete_goes_on_after_a_tenant_refuses_one(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, sp_enabled: Any, tmp_path: Path
) -> None:
    _delete(monkeypatch, sp_enabled)
    changes: list[httpx2.Request] = []
    apps = {
        TENANT_A: [
            _sp("sp-a1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True),
            _sp("sp-a2", "Reporter", REPORTER_APP_ID, enabled=True),
        ],
        TENANT_B: [_sp("sp-b1", "Contoso Cloud Backup", BACKUP_APP_ID, enabled=True)],
    }
    transport = _sp_services(apps, changes)

    def handler(request: httpx2.Request) -> httpx2.Response:
        # Alpha refuses to delete its first match, as Graph does an app it will not let go.
        if request.method == "DELETE" and request.url.path.endswith("/sp-a1"):
            error = {"code": "Authorization_RequestDenied", "message": "Insufficient privileges."}
            return httpx2.Response(403, json={"error": error})
        return transport.handle_request(request)

    assert _run_all(sp_enabled, httpx2.MockTransport(handler)) == 1

    # The refusal ends Alpha's run, so its other match stays; beta corp's is still deleted.
    assert [(bearer_tenant(r), r.url.path) for r in changes] == [
        (TENANT_B, "/v1.0/servicePrincipals/sp-b1")
    ]
    assert _sheets(_output(tmp_path))["Tenants"][1:] == [
        ("Alpha", "error", TENANT_A, None),
        ("beta corp", "ok", TENANT_B, "deleted"),
    ]


# --- entra_sign_in_logs.py ---------------------------------------------------------------

SIGN_IN_FILTER = "userPrincipalName eq 'admin@contoso.com'"


def _sign_in(number: int) -> dict[str, Any]:
    """Build a sign-in as Graph returns it, in part."""
    return {
        "id": f"si-{number}",
        "userPrincipalName": "admin@contoso.com",
        "status": {"errorCode": 0},
    }


def _sign_in_services(
    sign_ins: dict[str, list[dict[str, Any]]], reads: list[httpx2.Request], version: str = "beta"
) -> httpx2.MockTransport:
    """Fake Graph holding *sign_ins* per tenant under *version*, recording every read of them.

    A tenant's sign-ins come one per page, each page but the last with a next link, so a read
    of more than one follows the links.
    """

    def handler(request: httpx2.Request) -> httpx2.Response:
        held = sign_ins.get(bearer_tenant(request), [])
        if request.url.path == f"/{version}/auditLogs/signIns":
            reads.append(request)
            index = int(request.url.params.get("$skiptoken", "0"))
            page: dict[str, Any] = {"value": held[index : index + 1]}
            if index + 1 < len(held):
                page["@odata.nextLink"] = (
                    f"https://graph.microsoft.com/{version}/auditLogs/signIns"
                    f"?$skiptoken={index + 1}"
                )
            return httpx2.Response(200, json=page)
        return fake_services(request)

    return httpx2.MockTransport(handler)


@pytest.fixture
def sign_in_logs(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """The example, run through ``main()`` with ``--include all``, writing to a temp dir.

    Its filter is :data:`SIGN_IN_FILTER`, as the command line would set it.
    """
    monkeypatch.setenv(ENV_PREFIX + "PARTNER_TENANT_ID", PARTNER_TENANT)
    monkeypatch.setenv(ENV_PREFIX + "USERNAME", USER)
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)
    module = load_script(ENTRA_SIGN_IN_LOGS)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path / "reports")
    monkeypatch.setattr(module, "FILTER", SIGN_IN_FILTER)
    return module


def _sign_in_output(tmp_path: Path, tenant: str | None = None) -> Path:
    """Return where the example writes today's workbook, or with *tenant* that tenant's JSON."""
    today = datetime.date.today().isoformat()
    if tenant is None:
        return tmp_path / "reports" / f"EntraSignInLogs_{today}.xlsx"
    return tmp_path / "reports" / f"EntraSignInLogs_{tenant}_{today}.json"


def test_sign_in_logs_writes_json_per_tenant_with_matches_and_a_count_workbook(
    fake_msal: FakeMsal,
    sign_in_logs: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    reads: list[httpx2.Request] = []
    held = [_sign_in(1), _sign_in(2), _sign_in(3)]

    assert _run_all(sign_in_logs, _sign_in_services({TENANT_A: held}, reads)) == 0

    alpha_json = _sign_in_output(tmp_path, "Alpha")
    workbook = _sign_in_output(tmp_path)
    assert sorted(p.name for p in (tmp_path / "reports").iterdir()) == sorted(
        [alpha_json.name, workbook.name]
    )
    # Every page, whole, as Graph returned it.
    assert json.loads(alpha_json.read_text(encoding="utf-8")) == held
    assert capsys.readouterr().out == f"Wrote {workbook}\n"
    assert _sheets(workbook) == {
        "Tenants": [
            (SIGN_IN_FILTER, None, None, None, None),
            ("Tenant", "TenantId", "Status", "SignIns", "Error"),
            ("Alpha", TENANT_A, "ok", 3, None),
            ("beta corp", TENANT_B, "ok", 0, None),
        ]
    }
    # The filter goes on the first page; next links carry their own query.
    firsts = [r for r in reads if "$skiptoken" not in r.url.params]
    assert [(bearer_tenant(r), r.url.params["$filter"]) for r in firsts] == [
        (TENANT_A, SIGN_IN_FILTER),
        (TENANT_B, SIGN_IN_FILTER),
    ]
    assert len(reads) == 4
    # Beta unless --v1 says otherwise.
    assert {(r.url.host, r.url.path) for r in reads} == {
        ("graph.microsoft.com", "/beta/auditLogs/signIns")
    }
    assert sorted(
        r.getMessage() for r in caplog.records if r.filename == ENTRA_SIGN_IN_LOGS.name
    ) == [f"3 sign-ins match; wrote {alpha_json}", "No sign-ins match"]


def test_sign_in_logs_v1_reads_the_v1_endpoint(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, sign_in_logs: Any, tmp_path: Path
) -> None:
    monkeypatch.setattr(sign_in_logs, "V1", True)
    reads: list[httpx2.Request] = []
    transport = _sign_in_services({TENANT_B: [_sign_in(1)]}, reads, version="v1.0")

    assert _run_all(sign_in_logs, transport) == 0

    assert [(bearer_tenant(r), r.url.host, r.url.path) for r in reads] == [
        (TENANT_A, "graph.microsoft.com", "/v1.0/auditLogs/signIns"),
        (TENANT_B, "graph.microsoft.com", "/v1.0/auditLogs/signIns"),
    ]
    beta_json = _sign_in_output(tmp_path, "beta_corp")
    assert json.loads(beta_json.read_text(encoding="utf-8")) == [_sign_in(1)]
    rows = _sheets(_sign_in_output(tmp_path))["Tenants"]
    assert rows[0][0] == SIGN_IN_FILTER
    assert [row[3] for row in rows[2:]] == [0, 1]


def test_sign_in_logs_report_goes_on_after_a_tenant_error(
    fake_msal: FakeMsal, sign_in_logs: Any, tmp_path: Path
) -> None:
    reads: list[httpx2.Request] = []
    held = _sign_in_services({TENANT_B: [_sign_in(1)]}, reads)

    def handler(request: httpx2.Request) -> httpx2.Response:
        if request.url.path == "/beta/auditLogs/signIns" and bearer_tenant(request) == TENANT_A:
            error = {
                "code": "Authentication_RequestFromNonPremiumTenantOrB2CTenant",
                "message": "Neither tenant is B2C or tenant doesn't have premium license",
            }
            return httpx2.Response(403, json={"error": error})
        return held.handle_request(request)

    assert _run_all(sign_in_logs, httpx2.MockTransport(handler)) == 1

    alpha, beta = _sheets(_sign_in_output(tmp_path))["Tenants"][2:]
    assert alpha[:4] == ("Alpha", TENANT_A, "error", None)
    assert "premium license" in alpha[4]
    assert beta == ("beta corp", TENANT_B, "ok", 1, None)


def test_sign_in_logs_workbook_has_the_filter_title_and_a_table_below_it(
    sign_in_logs: Any, tmp_path: Path
) -> None:
    sign_in_logs.report(
        [
            TenantResult(Tenant(TENANT_A, "Alpha"), "ok", 2),
            TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
        ]
    )
    sheet = load_workbook(_sign_in_output(tmp_path))["Tenants"]
    assert sheet["A1"].value == SIGN_IN_FILTER
    assert sheet["A1"].font.b
    assert (sheet["A1"].font.sz, sheet["B3"].font.sz) == (20, 11)
    assert {cell.font.name for row in sheet.iter_rows() for cell in row} == {"Consolas"}
    # A theme scheme would make Excel show the theme's font, Calibri, in Consolas's place.
    assert {cell.font.scheme for row in sheet.iter_rows() for cell in row} == {None}
    assert sheet.column_dimensions["B"].width >= len(TENANT_A) + 2
    # The title is in A1 alone, and runs on across B1:E1, which stay empty.
    assert not sheet.merged_cells.ranges
    assert [cell.value for cell in sheet[1][1:]] == [None] * 4
    assert list(sheet.iter_rows(min_row=3, values_only=True)) == [
        ("Alpha", TENANT_A, "ok", 2, None),
        ("beta corp", TENANT_B, "auth_failed", None, "AADSTS65001"),
    ]
    assert sheet.freeze_panes == "A3"
    (table,) = sheet.tables.values()
    assert table.ref == "A2:E4"
    assert table.tableStyleInfo is not None
    assert table.tableStyleInfo.name == "TableStyleMedium8"
    # The workbook's default font: a cell nothing was written to has it too. Reading it
    # creates the cell, so this comes last.
    assert sheet["Z99"].font.name == "Consolas"


def test_sign_in_logs_writes_nothing_when_no_tenant_ran(
    sign_in_logs: Any, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    sign_in_logs.report([])
    assert not (tmp_path / "reports").exists()
    assert capsys.readouterr().out == "No tenants ran; no workbook written.\n"


def test_sign_in_logs_keeps_an_earlier_runs_files(
    fake_msal: FakeMsal,
    sign_in_logs: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    earlier_json = _sign_in_output(tmp_path, "Alpha")
    earlier_workbook = _sign_in_output(tmp_path)
    earlier_json.parent.mkdir(parents=True)
    earlier_json.write_bytes(b"earlier sign-ins")
    earlier_workbook.write_bytes(b"earlier workbook")

    transport = _sign_in_services({TENANT_A: [_sign_in(1)]}, [])
    assert _run_all(sign_in_logs, transport) == 0

    assert earlier_json.read_bytes() == b"earlier sign-ins"
    assert earlier_workbook.read_bytes() == b"earlier workbook"
    new_json = earlier_json.with_name(f"{earlier_json.stem}_2.json")
    new_workbook = earlier_workbook.with_name(f"{earlier_workbook.stem}_2.xlsx")
    assert json.loads(new_json.read_text(encoding="utf-8")) == [_sign_in(1)]
    assert _sheets(new_workbook)["Tenants"][0][0] == SIGN_IN_FILTER
    assert capsys.readouterr().out == f"Wrote {new_workbook}\n"
    assert f"1 sign-ins match; wrote {new_json}" in caplog.messages


def test_sign_in_logs_numbers_a_file_past_every_name_taken(
    sign_in_logs: Any, tmp_path: Path
) -> None:
    out = tmp_path / "reports"
    out.mkdir()
    for name in ("Alpha.json", "Alpha_2.json", "Alpha_4.json"):
        (out / name).write_bytes(name.encode())

    paths = []
    for _ in range(3):
        path, file = sign_in_logs._create("Alpha", ".json")
        with file:
            file.write(b"new")
        paths.append(path.name)

    assert paths == ["Alpha_3.json", "Alpha_5.json", "Alpha_6.json"]
    for name in ("Alpha.json", "Alpha_2.json", "Alpha_4.json"):
        assert (out / name).read_bytes() == name.encode()


# --- entra_audit_logs.py -----------------------------------------------------------------

AUDIT_FILTER = "activityDisplayName eq 'Add member to role'"


def _directory_audit(number: int) -> dict[str, Any]:
    """Build an audit event as Graph returns it, in part."""
    return {
        "id": f"da-{number}",
        "activityDisplayName": "Add member to role",
        "result": "success",
    }


def _directory_audit_services(
    events: dict[str, list[dict[str, Any]]], reads: list[httpx2.Request], version: str = "beta"
) -> httpx2.MockTransport:
    """Fake Graph holding *events* per tenant under *version*, recording every read of them.

    A tenant's audit events come one per page, each page but the last with a next link, so a
    read of more than one follows the links.
    """

    def handler(request: httpx2.Request) -> httpx2.Response:
        held = events.get(bearer_tenant(request), [])
        if request.url.path == f"/{version}/auditLogs/directoryAudits":
            reads.append(request)
            index = int(request.url.params.get("$skiptoken", "0"))
            page: dict[str, Any] = {"value": held[index : index + 1]}
            if index + 1 < len(held):
                page["@odata.nextLink"] = (
                    f"https://graph.microsoft.com/{version}/auditLogs/directoryAudits"
                    f"?$skiptoken={index + 1}"
                )
            return httpx2.Response(200, json=page)
        return fake_services(request)

    return httpx2.MockTransport(handler)


@pytest.fixture
def audit_logs(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """The example, run through ``main()`` with ``--include all``, writing to a temp dir.

    Its filter is :data:`AUDIT_FILTER`, as the command line would set it.
    """
    monkeypatch.setenv(ENV_PREFIX + "PARTNER_TENANT_ID", PARTNER_TENANT)
    monkeypatch.setenv(ENV_PREFIX + "USERNAME", USER)
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)
    module = load_script(ENTRA_AUDIT_LOGS)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path / "reports")
    monkeypatch.setattr(module, "FILTER", AUDIT_FILTER)
    return module


def _audit_logs_output(tmp_path: Path, tenant: str | None = None) -> Path:
    """Return where the example writes today's workbook, or with *tenant* that tenant's JSON."""
    today = datetime.date.today().isoformat()
    if tenant is None:
        return tmp_path / "reports" / f"EntraAuditLogs_{today}.xlsx"
    return tmp_path / "reports" / f"EntraAuditLogs_{tenant}_{today}.json"


def test_audit_logs_writes_json_per_tenant_with_matches_and_a_count_workbook(
    fake_msal: FakeMsal,
    audit_logs: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    reads: list[httpx2.Request] = []
    held = [_directory_audit(1), _directory_audit(2), _directory_audit(3)]

    assert _run_all(audit_logs, _directory_audit_services({TENANT_A: held}, reads)) == 0

    alpha_json = _audit_logs_output(tmp_path, "Alpha")
    workbook = _audit_logs_output(tmp_path)
    assert sorted(p.name for p in (tmp_path / "reports").iterdir()) == sorted(
        [alpha_json.name, workbook.name]
    )
    # Every page, whole, as Graph returned it.
    assert json.loads(alpha_json.read_text(encoding="utf-8")) == held
    assert capsys.readouterr().out == f"Wrote {workbook}\n"
    assert _sheets(workbook) == {
        "Tenants": [
            (AUDIT_FILTER, None, None, None, None),
            ("Tenant", "TenantId", "Status", "AuditEvents", "Error"),
            ("Alpha", TENANT_A, "ok", 3, None),
            ("beta corp", TENANT_B, "ok", 0, None),
        ]
    }
    # The filter goes on the first page; next links carry their own query.
    firsts = [r for r in reads if "$skiptoken" not in r.url.params]
    assert [(bearer_tenant(r), r.url.params["$filter"]) for r in firsts] == [
        (TENANT_A, AUDIT_FILTER),
        (TENANT_B, AUDIT_FILTER),
    ]
    assert len(reads) == 4
    # Beta unless --v1 says otherwise.
    assert {(r.url.host, r.url.path) for r in reads} == {
        ("graph.microsoft.com", "/beta/auditLogs/directoryAudits")
    }
    assert sorted(
        r.getMessage() for r in caplog.records if r.filename == ENTRA_AUDIT_LOGS.name
    ) == [f"3 audit events match; wrote {alpha_json}", "No audit events match"]


def test_audit_logs_v1_reads_the_v1_endpoint(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, audit_logs: Any, tmp_path: Path
) -> None:
    monkeypatch.setattr(audit_logs, "V1", True)
    reads: list[httpx2.Request] = []
    transport = _directory_audit_services({TENANT_B: [_directory_audit(1)]}, reads, "v1.0")

    assert _run_all(audit_logs, transport) == 0

    assert [(bearer_tenant(r), r.url.host, r.url.path) for r in reads] == [
        (TENANT_A, "graph.microsoft.com", "/v1.0/auditLogs/directoryAudits"),
        (TENANT_B, "graph.microsoft.com", "/v1.0/auditLogs/directoryAudits"),
    ]
    beta_json = _audit_logs_output(tmp_path, "beta_corp")
    assert json.loads(beta_json.read_text(encoding="utf-8")) == [_directory_audit(1)]
    rows = _sheets(_audit_logs_output(tmp_path))["Tenants"]
    assert rows[0][0] == AUDIT_FILTER
    assert [row[3] for row in rows[2:]] == [0, 1]


def test_audit_logs_report_goes_on_after_a_tenant_error(
    fake_msal: FakeMsal, audit_logs: Any, tmp_path: Path
) -> None:
    reads: list[httpx2.Request] = []
    held = _directory_audit_services({TENANT_B: [_directory_audit(1)]}, reads)

    def handler(request: httpx2.Request) -> httpx2.Response:
        path = "/beta/auditLogs/directoryAudits"
        if request.url.path == path and bearer_tenant(request) == TENANT_A:
            error = {
                "code": "Authentication_RequestFromUnsupportedUserRole",
                "message": "User is not in the allowed roles",
            }
            return httpx2.Response(403, json={"error": error})
        return held.handle_request(request)

    assert _run_all(audit_logs, httpx2.MockTransport(handler)) == 1

    alpha, beta = _sheets(_audit_logs_output(tmp_path))["Tenants"][2:]
    assert alpha[:4] == ("Alpha", TENANT_A, "error", None)
    assert "allowed roles" in alpha[4]
    assert beta == ("beta corp", TENANT_B, "ok", 1, None)


def test_audit_logs_workbook_has_the_filter_title_and_a_table_below_it(
    audit_logs: Any, tmp_path: Path
) -> None:
    audit_logs.report(
        [
            TenantResult(Tenant(TENANT_A, "Alpha"), "ok", 2),
            TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
        ]
    )
    sheet = load_workbook(_audit_logs_output(tmp_path))["Tenants"]
    assert sheet["A1"].value == AUDIT_FILTER
    assert sheet["A1"].font.b
    assert (sheet["A1"].font.sz, sheet["B3"].font.sz) == (20, 11)
    # The title is in A1 alone, and runs on across B1:E1, which stay empty.
    assert not sheet.merged_cells.ranges
    assert [cell.value for cell in sheet[1][1:]] == [None] * 4
    assert list(sheet.iter_rows(min_row=3, values_only=True)) == [
        ("Alpha", TENANT_A, "ok", 2, None),
        ("beta corp", TENANT_B, "auth_failed", None, "AADSTS65001"),
    ]
    assert sheet.freeze_panes == "A3"
    (table,) = sheet.tables.values()
    assert table.ref == "A2:E4"
    assert table.tableStyleInfo is not None
    assert table.tableStyleInfo.name == "TableStyleMedium8"
    _assert_consolas(_audit_logs_output(tmp_path))


def test_audit_logs_writes_nothing_when_no_tenant_ran(
    audit_logs: Any, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    audit_logs.report([])
    assert not (tmp_path / "reports").exists()
    assert capsys.readouterr().out == "No tenants ran; no workbook written.\n"


def test_audit_logs_keeps_an_earlier_runs_files(
    fake_msal: FakeMsal,
    audit_logs: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    earlier_json = _audit_logs_output(tmp_path, "Alpha")
    earlier_workbook = _audit_logs_output(tmp_path)
    earlier_json.parent.mkdir(parents=True)
    earlier_json.write_bytes(b"earlier audit events")
    earlier_workbook.write_bytes(b"earlier workbook")

    transport = _directory_audit_services({TENANT_A: [_directory_audit(1)]}, [])
    assert _run_all(audit_logs, transport) == 0

    assert earlier_json.read_bytes() == b"earlier audit events"
    assert earlier_workbook.read_bytes() == b"earlier workbook"
    new_json = earlier_json.with_name(f"{earlier_json.stem}_2.json")
    new_workbook = earlier_workbook.with_name(f"{earlier_workbook.stem}_2.xlsx")
    assert json.loads(new_json.read_text(encoding="utf-8")) == [_directory_audit(1)]
    assert _sheets(new_workbook)["Tenants"][0][0] == AUDIT_FILTER
    assert capsys.readouterr().out == f"Wrote {new_workbook}\n"
    assert f"1 audit events match; wrote {new_json}" in caplog.messages
    _assert_consolas(new_workbook)


# --- exchange_message_trace.py -----------------------------------------------------------

TRACE_V2 = "Get-MessageTraceV2"
TRACE_V1 = "Get-MessageTrace"
TRACE_PARAMETERS = {"SenderAddress": "admin@contoso.com", "Status": ["Failed", "Quarantined"]}
TRACE_FILTER = json.dumps(TRACE_PARAMETERS)


def _trace_record(number: int) -> dict[str, Any]:
    """Build a record as the cmdlet returns it, in part."""
    return {
        "MessageTraceId": f"mt-{number}",
        "Received@data.type": "System.DateTime",
        "Received": f"2026-10-28T12:00:{number:02}.0000000Z",
        "RecipientAddress": f"user{number}@contoso.com",
        "Status": "Failed",
    }


def _trace_domain(tenant: str) -> str:
    """Return the ``m365_domain`` the example's tenants file gives *tenant*."""
    return f"{tenant}.example"


def _trace_services(
    records: dict[str, list[dict[str, Any]]],
    queries: list[tuple[str, str, dict[str, Any]]],
    *,
    refusals: dict[tuple[str, str], httpx2.Response] | None = None,
) -> httpx2.MockTransport:
    """Fake Exchange holding *records* per tenant, recording each query's cmdlet and parameters.

    *refusals* holds the answer a tenant gives a cmdlet in place of its records, by the tenant and
    the cmdlet. :func:`_unknown_cmdlet` is how it answers one it does not know.
    """

    def handler(request: httpx2.Request) -> httpx2.Response:
        tenant = bearer_tenant(request)
        if request.url.path == f"/adminapi/beta/{tenant}/InvokeCommand":
            cmdlet = json.loads(request.content)["CmdletInput"]
            name = cmdlet["CmdletName"]
            queries.append((tenant, name, cmdlet["Parameters"]))
            refusal = (refusals or {}).get((tenant, name))
            if refusal is not None:
                return refusal
            return httpx2.Response(200, json={"value": records.get(tenant, [])})
        return fake_services(request)

    return httpx2.MockTransport(handler)


def _unknown_cmdlet() -> httpx2.Response:
    """Answer as Exchange does a cmdlet it does not know: a 403 with no body."""
    return httpx2.Response(403)


def _cmdlet_error(reason: str) -> httpx2.Response:
    """Answer as Exchange does a cmdlet that failed, with *reason* in its details."""
    error = {
        "code": "BadRequest",
        "message": "Error executing cmdlet",
        "details": [{"message": f"Ex1|Exception|{reason}"}],
    }
    return httpx2.Response(400, json={"error": error})


@pytest.fixture
def message_trace(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """The example, run through ``main()`` with ``--include all``, writing to a temp dir.

    Its filter is :data:`TRACE_FILTER`, as the command line would set it, and a tenants file
    gives each tenant its domain.
    """
    tenants_file = tmp_path / "tenants.csv"
    tenants_file.write_text(
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url\n"
        f"{TENANT_A},Alpha,{_trace_domain(TENANT_A)},Commercial,,TRUE,\n"
        f"{TENANT_B},beta corp,{_trace_domain(TENANT_B)},Commercial,,TRUE,\n",
        encoding="utf-8",
    )
    monkeypatch.setenv(ENV_PREFIX + "TENANTS_FILE", str(tenants_file))
    monkeypatch.setenv(ENV_PREFIX + "PARTNER_TENANT_ID", PARTNER_TENANT)
    monkeypatch.setenv(ENV_PREFIX + "USERNAME", USER)
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)
    module = load_script(MESSAGE_TRACE)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path / "reports")
    monkeypatch.setattr(module, "FILTER", TRACE_FILTER)
    return module


def _trace_output(tmp_path: Path, tenant: str | None = None) -> Path:
    """Return where the example writes today's workbook, or with *tenant* that tenant's JSON."""
    today = datetime.date.today().isoformat()
    if tenant is None:
        return tmp_path / "reports" / f"ExchangeMessageTrace_{today}.xlsx"
    return tmp_path / "reports" / f"ExchangeMessageTrace_{tenant}_{today}.json"


def _trace_records_output(tmp_path: Path) -> Path:
    """Return where the example writes today's records workbook."""
    today = datetime.date.today().isoformat()
    return tmp_path / "reports" / f"ExchangeMessageTraceRecords_{today}.xlsx"


def test_message_trace_sends_the_filter_and_the_size_and_writes_json_and_a_count_workbook(
    fake_msal: FakeMsal,
    message_trace: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    queries: list[tuple[str, str, dict[str, Any]]] = []
    held = [_trace_record(1), _trace_record(2), _trace_record(3)]

    assert _run_all(message_trace, _trace_services({TENANT_A: held}, queries)) == 0

    # One query per tenant: the parameters given, and the size.
    sent = {**TRACE_PARAMETERS, "ResultSize": 5000}
    assert queries == [(TENANT_A, TRACE_V2, sent), (TENANT_B, TRACE_V2, sent)]
    alpha_json = _trace_output(tmp_path, "Alpha")
    workbook = _trace_output(tmp_path)
    records_workbook = _trace_records_output(tmp_path)
    assert sorted(p.name for p in (tmp_path / "reports").iterdir()) == sorted(
        [alpha_json.name, workbook.name, records_workbook.name]
    )
    # Every record, whole, as the cmdlet returned it.
    assert json.loads(alpha_json.read_text(encoding="utf-8")) == held
    assert capsys.readouterr().out == f"Wrote {workbook}\nWrote {records_workbook}\n"
    assert _sheets(workbook) == {
        "Tenants": [
            (TRACE_FILTER, None, None, None, None),
            ("Tenant", "TenantId", "Status", "Records", "Error"),
            ("Alpha", TENANT_A, "ok", 3, None),
            ("beta corp", TENANT_B, "ok", 0, None),
        ]
    }
    assert sorted(r.getMessage() for r in caplog.records if r.filename == MESSAGE_TRACE.name) == [
        f"3 records match; wrote {alpha_json}",
        "No records match",
    ]


def test_message_trace_empty_filter_sends_only_the_size(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, message_trace: Any
) -> None:
    monkeypatch.setattr(message_trace, "FILTER", "{}")
    monkeypatch.setattr(message_trace, "SIZE", 100)
    queries: list[tuple[str, str, dict[str, Any]]] = []

    assert _run_all(message_trace, _trace_services({}, queries)) == 0

    sent = {"ResultSize": 100}
    assert queries == [(TENANT_A, TRACE_V2, sent), (TENANT_B, TRACE_V2, sent)]


def test_message_trace_filters_own_size_wins(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, message_trace: Any
) -> None:
    monkeypatch.setattr(message_trace, "FILTER", '{"ResultSize": 7}')
    queries: list[tuple[str, str, dict[str, Any]]] = []

    assert _run_all(message_trace, _trace_services({}, queries)) == 0

    assert {tuple(parameters.items()) for _, _, parameters in queries} == {(("ResultSize", 7),)}


def test_message_trace_falls_back_to_the_older_cmdlet_on_a_403(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    message_trace: Any,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    monkeypatch.setattr(message_trace, "SIZE", 100)
    queries: list[tuple[str, str, dict[str, Any]]] = []
    held = [_trace_record(1), _trace_record(2)]
    refusals = {(TENANT_A, TRACE_V2): _unknown_cmdlet()}

    transport = _trace_services({TENANT_A: held}, queries, refusals=refusals)
    assert _run_all(message_trace, transport) == 0

    # The same filter, with the size in the older cmdlet's spelling; the other tenant is untouched.
    assert queries == [
        (TENANT_A, TRACE_V2, {**TRACE_PARAMETERS, "ResultSize": 100}),
        (TENANT_A, TRACE_V1, {**TRACE_PARAMETERS, "PageSize": 100}),
        (TENANT_B, TRACE_V2, {**TRACE_PARAMETERS, "ResultSize": 100}),
    ]
    assert json.loads(_trace_output(tmp_path, "Alpha").read_text(encoding="utf-8")) == held
    assert _sheets(_trace_output(tmp_path))["Tenants"][2:] == [
        ("Alpha", TENANT_A, "ok", 2, None),
        ("beta corp", TENANT_B, "ok", 0, None),
    ]
    assert "Exchange refused Get-MessageTraceV2; trying Get-MessageTrace" in caplog.messages


@pytest.mark.parametrize(
    ("older", "reason"),
    [
        (
            _cmdlet_error("Get-MessageTrace will start deprecating on September 1st, 2025."),
            "Get-MessageTrace failed: Get-MessageTrace will start deprecating",
        ),
        (_unknown_cmdlet(), "Get-MessageTrace failed: Exchange sent no error body"),
    ],
    ids=["reason given", "no reason given"],
)
def test_message_trace_error_names_both_cmdlets_when_both_are_refused(
    fake_msal: FakeMsal,
    message_trace: Any,
    tmp_path: Path,
    older: httpx2.Response,
    reason: str,
) -> None:
    refusals = {(TENANT_A, TRACE_V2): _unknown_cmdlet(), (TENANT_A, TRACE_V1): older}

    assert _run_all(message_trace, _trace_services({}, [], refusals=refusals)) == 1

    alpha, beta = _sheets(_trace_output(tmp_path))["Tenants"][2:]
    assert alpha[:4] == ("Alpha", TENANT_A, "error", None)
    assert "Exchange refused Get-MessageTraceV2 with a 403" in alpha[4]
    assert reason in alpha[4]
    assert beta == ("beta corp", TENANT_B, "ok", 0, None)


@pytest.mark.parametrize(
    "refusal",
    [httpx2.Response(401), _cmdlet_error("The interval is too long.")],
    ids=["401", "400"],
)
def test_message_trace_does_not_fall_back_on_any_other_error(
    fake_msal: FakeMsal, message_trace: Any, tmp_path: Path, refusal: httpx2.Response
) -> None:
    queries: list[tuple[str, str, dict[str, Any]]] = []
    refusals = {(TENANT_A, TRACE_V2): refusal}

    assert _run_all(message_trace, _trace_services({}, queries, refusals=refusals)) == 1

    assert TRACE_V1 not in {cmdlet for _, cmdlet, _ in queries}
    alpha, _ = _sheets(_trace_output(tmp_path))["Tenants"][2:]
    assert alpha[:3] == ("Alpha", TENANT_A, "error")
    assert "Exchange refused" not in alpha[4]


def test_message_trace_reports_a_filter_that_is_not_json(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, message_trace: Any, tmp_path: Path
) -> None:
    monkeypatch.setattr(message_trace, "FILTER", "not json")
    queries: list[tuple[str, str, dict[str, Any]]] = []

    assert _run_all(message_trace, _trace_services({}, queries)) == 1

    assert not queries
    alpha, beta = _sheets(_trace_output(tmp_path))["Tenants"][2:]
    assert alpha[:4] == ("Alpha", TENANT_A, "error", None)
    assert beta[:4] == ("beta corp", TENANT_B, "error", None)
    assert "Expecting value" in alpha[4]


def test_message_trace_report_goes_on_after_a_tenant_error(
    fake_msal: FakeMsal, message_trace: Any, tmp_path: Path
) -> None:
    reason = "The interval between StartDate and EndDate is too long."
    refusals = {(TENANT_A, TRACE_V2): _cmdlet_error(reason)}
    transport = _trace_services({TENANT_B: [_trace_record(1)]}, [], refusals=refusals)

    assert _run_all(message_trace, transport) == 1

    alpha, beta = _sheets(_trace_output(tmp_path))["Tenants"][2:]
    assert alpha[:4] == ("Alpha", TENANT_A, "error", None)
    assert "The interval between StartDate and EndDate is too long." in alpha[4]
    assert beta == ("beta corp", TENANT_B, "ok", 1, None)
    assert not _trace_output(tmp_path, "Alpha").exists()


def test_message_trace_records_workbook_lists_every_record_under_its_tenant_and_domain(
    fake_msal: FakeMsal, message_trace: Any, tmp_path: Path
) -> None:
    other = {"MessageTraceId": "mt-3", "FromIP": "203.0.113.10", "Subject": "=1+1"}
    held = {TENANT_A: [_trace_record(1), _trace_record(2)], TENANT_B: [other]}

    assert _run_all(message_trace, _trace_services(held, [])) == 0

    path = _trace_records_output(tmp_path)
    # Every property any record has, in the order they first appear, without the type notes.
    assert _sheets(path) == {
        "Records": [
            (TRACE_FILTER, *[None] * 7),
            (
                "Tenant",
                "Domain",
                "MessageTraceId",
                "Received",
                "RecipientAddress",
                "Status",
                "FromIP",
                "Subject",
            ),
            (
                "Alpha",
                _trace_domain(TENANT_A),
                "mt-1",
                _trace_record(1)["Received"],
                "user1@contoso.com",
                "Failed",
                None,
                None,
            ),
            (
                "Alpha",
                _trace_domain(TENANT_A),
                "mt-2",
                _trace_record(2)["Received"],
                "user2@contoso.com",
                "Failed",
                None,
                None,
            ),
            (
                "beta corp",
                _trace_domain(TENANT_B),
                "mt-3",
                None,
                None,
                None,
                "203.0.113.10",
                "=1+1",
            ),
        ]
    }
    sheet = load_workbook(path)["Records"]
    # A subject that looks like a formula stays text.
    assert sheet["H5"].data_type == "s"
    assert sheet["A1"].font.b
    assert sheet.freeze_panes == "A3"
    (table,) = sheet.tables.values()
    assert table.ref == "A2:H5"
    assert table.tableStyleInfo is not None
    assert table.tableStyleInfo.name == "TableStyleMedium8"
    _assert_consolas(path)


def test_message_trace_records_workbook_leaves_out_failed_tenants_and_writes_lists_as_json(
    message_trace: Any, tmp_path: Path
) -> None:
    record = {"MessageTraceId": "mt-1", "Recipients": ["a@contoso.com", "b@contoso.com"]}
    message_trace.report(
        [
            TenantResult(Tenant(TENANT_A, "Alpha", domain="a.example"), "ok", [record]),
            TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
        ]
    )
    assert _sheets(_trace_records_output(tmp_path))["Records"][2:] == [
        ("Alpha", "a.example", "mt-1", '["a@contoso.com", "b@contoso.com"]'),
    ]


def test_message_trace_leaves_the_domain_blank_where_the_tenants_file_has_none(
    message_trace: Any, tmp_path: Path
) -> None:
    message_trace.report([TenantResult(Tenant(TENANT_A, "Alpha"), "ok", [_trace_record(1)])])
    assert _sheets(_trace_records_output(tmp_path))["Records"][2][:3] == ("Alpha", None, "mt-1")


def test_message_trace_writes_no_records_workbook_when_nothing_matches(
    fake_msal: FakeMsal,
    message_trace: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    assert _run_all(message_trace, _trace_services({}, [])) == 0

    workbook = _trace_output(tmp_path)
    assert [p.name for p in (tmp_path / "reports").iterdir()] == [workbook.name]
    assert capsys.readouterr().out == (
        f"Wrote {workbook}\nNo records match; no records workbook written.\n"
    )


def test_message_trace_workbook_has_the_filter_title_and_a_table_below_it(
    message_trace: Any, tmp_path: Path
) -> None:
    message_trace.report(
        [
            TenantResult(Tenant(TENANT_A, "Alpha"), "ok", [_trace_record(1), _trace_record(2)]),
            TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
        ]
    )
    sheet = load_workbook(_trace_output(tmp_path))["Tenants"]
    assert sheet["A1"].value == TRACE_FILTER
    assert sheet["A1"].font.b
    assert (sheet["A1"].font.sz, sheet["B3"].font.sz) == (20, 11)
    # The title is in A1 alone, and runs on across B1:E1, which stay empty.
    assert not sheet.merged_cells.ranges
    assert [cell.value for cell in sheet[1][1:]] == [None] * 4
    assert list(sheet.iter_rows(min_row=3, values_only=True)) == [
        ("Alpha", TENANT_A, "ok", 2, None),
        ("beta corp", TENANT_B, "auth_failed", None, "AADSTS65001"),
    ]
    assert sheet.freeze_panes == "A3"
    (table,) = sheet.tables.values()
    assert table.ref == "A2:E4"
    assert table.tableStyleInfo is not None
    assert table.tableStyleInfo.name == "TableStyleMedium8"
    _assert_consolas(_trace_output(tmp_path))


def test_message_trace_writes_nothing_when_no_tenant_ran(
    message_trace: Any, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    message_trace.report([])
    assert not (tmp_path / "reports").exists()
    assert capsys.readouterr().out == "No tenants ran; no workbook written.\n"


def test_message_trace_keeps_an_earlier_runs_files(
    fake_msal: FakeMsal,
    message_trace: Any,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    earlier_json = _trace_output(tmp_path, "Alpha")
    earlier_workbook = _trace_output(tmp_path)
    earlier_records = _trace_records_output(tmp_path)
    earlier_json.parent.mkdir(parents=True)
    earlier_json.write_bytes(b"earlier records")
    earlier_workbook.write_bytes(b"earlier workbook")
    earlier_records.write_bytes(b"earlier records workbook")

    transport = _trace_services({TENANT_A: [_trace_record(1)]}, [])
    assert _run_all(message_trace, transport) == 0

    assert earlier_json.read_bytes() == b"earlier records"
    assert earlier_workbook.read_bytes() == b"earlier workbook"
    assert earlier_records.read_bytes() == b"earlier records workbook"
    new_json = earlier_json.with_name(f"{earlier_json.stem}_2.json")
    new_workbook = earlier_workbook.with_name(f"{earlier_workbook.stem}_2.xlsx")
    new_records = earlier_records.with_name(f"{earlier_records.stem}_2.xlsx")
    assert json.loads(new_json.read_text(encoding="utf-8")) == [_trace_record(1)]
    assert _sheets(new_workbook)["Tenants"][0][0] == TRACE_FILTER
    assert _sheets(new_records)["Records"][0][0] == TRACE_FILTER
    assert capsys.readouterr().out == f"Wrote {new_workbook}\nWrote {new_records}\n"
    assert f"1 records match; wrote {new_json}" in caplog.messages
    _assert_consolas(new_workbook)
    _assert_consolas(new_records)


# --- graph_scopes.py ---------------------------------------------------------------------


@pytest.fixture
def graph_scopes(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
    """The example, checking for two scopes and writing to a temp dir."""
    module = load_script(GRAPH_SCOPES)
    monkeypatch.setattr(module, "OUT_DIR", tmp_path / "reports")
    monkeypatch.setattr(module, "REQUESTED_SCOPES", ["User.Read.All", "Mail.Read"])
    return module


def _graph_scopes_results() -> list[TenantResult]:
    """Return one tenant whose app lacks Mail.Read, and one that failed."""
    held = {"appId": "app-1", "scopes": ["Policy.Read.All", "User.Read.All"]}
    return [
        TenantResult(Tenant(TENANT_A, "Alpha"), "ok", held),
        TenantResult(Tenant(TENANT_B, "beta corp"), "auth_failed", error="AADSTS65001"),
    ]


def test_graph_scopes_workbook_lists_every_tenant(
    graph_scopes: Any, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    graph_scopes.report(_graph_scopes_results())

    (path,) = (tmp_path / "reports").iterdir()
    assert path.name == f"GraphScopes_{datetime.date.today().isoformat()}.xlsx"
    assert capsys.readouterr().out == f"Wrote {path}\n"
    sheet = load_workbook(path)["Tenants"]
    assert list(sheet.iter_rows(values_only=True)) == [
        ("Tenant", "Status", "TenantId", "HasScopes", "Missing", "Scopes", "Error"),
        ("Alpha", "ok", TENANT_A, False, "Mail.Read", "Policy.Read.All User.Read.All", None),
        ("beta corp", "auth_failed", TENANT_B, None, None, None, "AADSTS65001"),
    ]
    assert sheet.freeze_panes == "A2"
    (table,) = sheet.tables.values()
    assert table.ref == "A1:G3"
    _assert_consolas(path)


def test_graph_scopes_keeps_an_earlier_runs_files(
    graph_scopes: Any,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr(graph_scopes, "EXPORT_JSON", True)
    today = datetime.date.today().isoformat()
    reports = tmp_path / "reports"
    reports.mkdir()
    earlier = [reports / f"GraphScopes_{today}.xlsx", reports / f"GraphScopes_{today}.json"]
    for path in earlier:
        path.write_bytes(path.name.encode())

    graph_scopes.report(_graph_scopes_results())

    for path in earlier:
        assert path.read_bytes() == path.name.encode()
    workbook, exported = (path.with_name(f"{path.stem}_2{path.suffix}") for path in earlier)
    assert capsys.readouterr().out == f"Wrote {workbook}\nWrote {exported}\n"
    assert graph_scopes.load_results(exported) == _graph_scopes_results()
    _assert_consolas(workbook)
