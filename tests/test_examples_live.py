"""Live tests for the examples: real tenants and real sign-ins, and they only read.

They run on the ``pilot`` group of the tenants file, the group the getting started guide has
you make to try a first script on, as the partner account in the config system.
"""

from __future__ import annotations

import argparse
import datetime
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from m365_gdap_runner import cli
from tests._gdap_fakes import load_script

pytestmark = [pytest.mark.e2e, pytest.mark.live, pytest.mark.usefixtures("settings")]

MESSAGE_TRACE = Path(__file__).resolve().parents[1] / "examples" / "exchange_message_trace.py"


def _window(days: int = 9) -> dict[str, str]:
    """Return a filter's dates, ending an hour ago, so that new mail cannot change the answer."""
    end = datetime.datetime.now(datetime.UTC) - datetime.timedelta(hours=1)
    start = end - datetime.timedelta(days=days)
    return {
        "StartDate": f"{start:%Y-%m-%dT%H:%M:%SZ}",
        "EndDate": f"{end:%Y-%m-%dT%H:%M:%SZ}",
    }


def _run_trace(
    monkeypatch: pytest.MonkeyPatch,
    out: Path,
    parameters: dict[str, Any],
    size: int | None = None,
) -> tuple[int, list[tuple[Any, ...]], list[dict[str, Any]]]:
    """Run the example on the pilot group.

    Args:
        monkeypatch: Keeps the run from stacking log handlers on the root logger.
        out: The folder to write to.
        parameters: The filter, before it is written as JSON.
        size: What ``--size`` would set, instead of the script's default.

    Returns:
        The exit code, the workbook's tenant rows, and every tenant's records together.
    """
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)
    module: Any = load_script(MESSAGE_TRACE)
    module.FILTER = json.dumps(parameters)
    module.OUT_DIR = out
    if size is not None:
        module.SIZE = size
    parser = argparse.ArgumentParser()
    cli.add_arguments(parser)
    code = cli.main(module, parser.parse_args(["--include", "pilot"]))
    (workbook,) = out.glob("*.xlsx")
    rows = list(load_workbook(workbook)["Tenants"].iter_rows(min_row=3, values_only=True))
    records = [
        record
        for path in sorted(out.glob("*.json"))
        for record in json.loads(path.read_text(encoding="utf-8"))
    ]
    return code, rows, records


def _keys(records: list[dict[str, Any]]) -> set[tuple[str, str]]:
    """Return what identifies each record: one message to one recipient."""
    return {(record["MessageTraceId"], record["RecipientAddress"]) for record in records}


def test_message_trace_writes_every_record_the_cmdlet_returns(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    code, rows, records = _run_trace(monkeypatch, tmp_path, _window())
    assert code == 0
    assert {row[2] for row in rows} == {"ok"}
    assert sum(row[3] for row in rows) == len(records)


def test_message_trace_empty_filter_runs(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    code, rows, _ = _run_trace(monkeypatch, tmp_path, {})
    assert code == 0
    assert {row[2] for row in rows} == {"ok"}


def test_message_trace_size_reaches_the_cmdlet(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    code, rows, _ = _run_trace(monkeypatch, tmp_path, _window(), size=1)
    assert code == 0
    assert {row[2] for row in rows} == {"ok"}
    assert max(row[3] for row in rows) <= 1


def test_message_trace_result_size_in_the_filter_reaches_the_cmdlet(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    code, rows, _ = _run_trace(monkeypatch, tmp_path, {"ResultSize": 1, **_window()})
    assert code == 0
    assert {row[2] for row in rows} == {"ok"}
    assert max(row[3] for row in rows) <= 1


@pytest.mark.parametrize(
    "parameters",
    [
        {"Status": ["Delivered", "Failed"]},
        {"SenderAddress": ["nobody@contoso.com", "nothing@contoso.com"]},
        {"Subject": "a b", "SubjectFilterType": "Contains"},
        {"status": "Delivered"},
    ],
    ids=["several statuses", "several senders", "subject", "lowercase name"],
)
def test_message_trace_filters_the_cmdlet_accepts(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, parameters: dict[str, Any]
) -> None:
    code, rows, _ = _run_trace(monkeypatch, tmp_path, {**parameters, **_window()})
    assert code == 0
    assert {row[2] for row in rows} == {"ok"}


def test_message_trace_several_statuses_return_only_those(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    filters = {"Status": ["Delivered", "Failed"], **_window()}
    _, _, records = _run_trace(monkeypatch, tmp_path, filters)
    assert {record["Status"] for record in records} <= {"Delivered", "Failed"}


def test_message_trace_reads_on_from_a_records_time_and_recipient(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    window = _window()
    _, _, first = _run_trace(monkeypatch, tmp_path / "first", {"ResultSize": 2, **window})
    if len(first) < 2:
        pytest.skip("The pilot tenants have too little mail in the window to read on from.")

    after = {
        "ResultSize": 2,
        "StartDate": window["StartDate"],
        "EndDate": first[-1]["Received"],
        "StartingRecipientAddress": first[-1]["RecipientAddress"],
    }
    code, _, second = _run_trace(monkeypatch, tmp_path / "second", after)

    assert code == 0
    assert not _keys(first) & _keys(second)


def test_message_trace_reports_what_exchange_refuses(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    code, rows, records = _run_trace(monkeypatch, tmp_path, _window(days=11))
    assert code == 1
    assert {row[2] for row in rows} == {"error"}
    assert all("can't be longer than 10 days" in row[4] for row in rows)
    assert records == []


def test_message_trace_reports_one_date_without_the_other(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    start = {"StartDate": _window()["StartDate"]}
    code, rows, records = _run_trace(monkeypatch, tmp_path, start)
    assert code == 1
    assert {row[2] for row in rows} == {"error"}
    assert records == []
