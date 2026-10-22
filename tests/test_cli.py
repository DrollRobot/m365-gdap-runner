"""Tests for main(), add_arguments() and the m365-gdap-runner-customer-export command in cli.py.

End-to-end runs use the fake msal module and a fake service transport.
"""

from __future__ import annotations

import argparse
import contextlib
import gzip
import json
import logging
import os
import re
import sys
import textwrap
from collections.abc import Callable, Iterator
from pathlib import Path
from types import ModuleType
from typing import Any

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from m365_gdap_runner import cli
from m365_gdap_runner.config.paths import CONFIG_DIR_ENV
from m365_gdap_runner.config.schema import ENV_PREFIX
from m365_gdap_runner.runner import TenantResult, run_payload
from m365_gdap_runner.session import PartnerSession
from m365_gdap_runner.tenants import OWN_COLUMNS, TenantFile
from tests._gdap_fakes import (
    CONSENT_REQUIRED,
    GAMMA_ADMIN,
    NOT_IN_TENANT,
    PARTNER_TENANT,
    TENANT_A,
    TENANT_B,
    TENANT_C,
    USER,
    FakeMsal,
    bearer_tenant,
    fake_services,
    load_script,
)

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"

# Microsoft Graph's own application ID: the resource of an app's Graph grant.
GRAPH_APP_ID = "00000003-0000-0000-c000-000000000000"
LICENSED_USERS = load_script(EXAMPLES / "licensed_users.py")
AUDIT_DISABLED = load_script(EXAMPLES / "audit_disabled.py")

# The real one: the fixture below replaces it so tests do not stack handlers on the root logger.
CONFIGURE_LOGGING = cli.configure_logging

# A script whose report() keeps what it is given, so a test can see every tenant's result.
RECORDER = """
    SCOPES = ["User.Read.All"]
    REPORTED = []

    def run(ctx):
        return [user["userPrincipalName"] for user in ctx.graph.get_all("/users")]

    def report(results):
        REPORTED.append(results)
"""

Handler = Callable[[httpx.Request], httpx.Response]


@pytest.fixture(autouse=True)
def _config(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Isolate config and supply the partner tenant and user through the environment."""
    for key in list(os.environ):
        if key.startswith(ENV_PREFIX):
            monkeypatch.delenv(key)
    monkeypatch.setenv(CONFIG_DIR_ENV, str(tmp_path / "config"))
    monkeypatch.setenv(ENV_PREFIX + "PARTNER_TENANT_ID", PARTNER_TENANT)
    monkeypatch.setenv(ENV_PREFIX + "USERNAME", USER)
    monkeypatch.setattr(cli, "configure_logging", lambda verbosity: None)


@pytest.fixture
def recorder(tmp_path: Path) -> ModuleType:
    return _script(tmp_path, RECORDER)


def _args(*argv: str) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    cli.add_arguments(parser)
    return parser.parse_args(list(argv))


def _main(script: ModuleType, *argv: str, handler: Handler = fake_services) -> int:
    return cli.main(script, _args(*argv), transport=httpx.MockTransport(handler))


def _list_tenants(*argv: str) -> int:
    return cli.list_tenants(list(argv), transport=httpx.MockTransport(fake_services))


def _script(tmp_path: Path, body: str, name: str = "script") -> ModuleType:
    path = tmp_path / f"{name}.py"
    path.write_text(textwrap.dedent(body), encoding="utf-8")
    return load_script(path)


def _reported(script: ModuleType) -> list[TenantResult]:
    """Return what the script's report() was given by the latest run."""
    reported: list[list[TenantResult]] = script.__dict__["REPORTED"]
    return reported[-1]


def _tenant_ids(script: ModuleType) -> list[str]:
    return [result.tenant.tenant_id for result in _reported(script)]


# --- arguments -------------------------------------------------------------------------


@pytest.mark.unit
def test_add_arguments_defaults() -> None:
    args = _args("--include", "all")
    assert (args.jobs, args.continue_on_error, args.grant_when_possible) == (None, None, None)
    assert (args.include, args.exclude) == (["all"], None)
    assert (args.partner_tenant_id, args.username, args.client_id, args.cache, args.cloud) == (
        None,
        None,
        None,
        None,
        None,
    )
    assert (args.tenants_file, args.profile, args.verbose) == (None, None, 0)


@pytest.mark.unit
@pytest.mark.parametrize(
    "argv",
    [
        ["--jobs", "0"],
        ["--jobs", "two"],
        ["--cache", "plaintext"],
        # The Windows broker is gone, and the flag that chose it with it.
        ["--sign-in", "broker"],
        ["--format", "json"],
    ],
    ids=str,
)
def test_the_parser_refuses_bad_values(argv: list[str]) -> None:
    with pytest.raises(SystemExit) as info:
        _args("--include", "all", *argv)
    assert info.value.code == 2


@pytest.mark.unit
def test_the_parser_main_builds_is_described_by_the_scripts_docstring() -> None:
    parser = cli._parser(LICENSED_USERS)
    assert parser.description == "Example script: licensed users in each tenant."
    assert "--include" in parser.format_help()


@pytest.mark.unit
def test_main_returns_130_on_interrupt(monkeypatch: pytest.MonkeyPatch) -> None:
    def interrupted(*args: object, **kwargs: object) -> int:
        raise KeyboardInterrupt

    monkeypatch.setattr(cli, "_run", interrupted)
    assert cli.main(LICENSED_USERS, _args("--include", "all")) == 130


@contextlib.contextmanager
def _restoring_loggers() -> Iterator[None]:
    """Undo what configure_logging() and the debug log do to the root and library loggers.

    Use it inside a test, not a fixture: pytest swaps its own root handlers between phases.
    """
    root = logging.getLogger()
    handlers = list(root.handlers)
    saved = {name: logging.getLogger(name).level for name in ("", *cli._LIBRARY_LOGGERS)}
    try:
        yield
    finally:
        for handler in root.handlers[len(handlers) :]:
            root.removeHandler(handler)
            handler.close()
        for name, level in saved.items():
            logging.getLogger(name).setLevel(level)


@pytest.mark.unit
def test_configure_logging_quiets_libraries_unless_very_verbose() -> None:
    root = logging.getLogger()
    with _restoring_loggers():
        CONFIGURE_LOGGING(0)
        assert root.level == logging.INFO
        assert logging.getLogger("azure_auth").level == logging.WARNING
        CONFIGURE_LOGGING(2)
        assert root.level == logging.DEBUG
        assert logging.getLogger("msal").level == logging.DEBUG


@pytest.mark.unit
@pytest.mark.parametrize(
    ("name", "level", "kept"),
    [
        ("httpx", logging.DEBUG, False),
        ("azure.identity", logging.INFO, False),
        ("httpx", logging.WARNING, True),
        ("httpx_extras", logging.INFO, True),
        ("m365_gdap_runner.session", logging.DEBUG, True),
    ],
)
def test_the_terminal_drops_library_records_below_warning(
    name: str, level: int, kept: bool
) -> None:
    record = logging.LogRecord(name, level, __file__, 1, "message", None, None)
    assert cli._quiet_libraries(record) is kept


@pytest.mark.integration
def test_settings_flags_override_config() -> None:
    args = _args(
        "--partner-tenant-id",
        "other.example",
        "--username",
        "someone@other.example",
        "--client-id",
        "custom",
        "--cache",
        "disk",
        "--cloud",
        "USGov",
        "--tenants-file",
        "C:/msp/tenants.csv",
        "--jobs",
        "3",
        "--include",
        "all",
    )
    settings = cli.settings_from_args(args)
    assert (
        settings.partner_tenant_id,
        settings.username,
        settings.client_id,
        settings.cache,
        settings.cloud,
        settings.tenants_file,
        settings.jobs,
    ) == (
        "other.example",
        "someone@other.example",
        "custom",
        "disk",
        "USGov",
        "C:/msp/tenants.csv",
        3,
    )


# --- how a script calls main() ---------------------------------------------------------


@pytest.mark.integration
@pytest.mark.acceptance
def test_main_reads_the_command_line_when_the_script_brings_no_parser(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, recorder: ModuleType
) -> None:
    monkeypatch.setattr(sys, "argv", ["script.py", "--include", "Alpha"])
    assert cli.main(recorder, transport=httpx.MockTransport(fake_services)) == 0
    assert _tenant_ids(recorder) == [TENANT_A]


@pytest.mark.integration
def test_main_finds_the_script_by_module_name(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, recorder: ModuleType
) -> None:
    monkeypatch.setitem(sys.modules, "__main__", recorder)
    assert (
        cli.main(
            "__main__", _args("--include", "all"), transport=httpx.MockTransport(fake_services)
        )
        == 0
    )
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_script_adds_flags_of_its_own(fake_msal: FakeMsal, tmp_path: Path) -> None:
    script = _script(
        tmp_path,
        """
        CLIENTS = ["exchange"]
        LABEL = "unset"
        REPORTED = []

        def run(ctx):
            return {"label": LABEL}

        def report(results):
            REPORTED.append(results)
        """,
    )
    parser = argparse.ArgumentParser()
    parser.add_argument("--label", default="none")
    cli.add_arguments(parser)
    assert "--label" in parser.format_help()
    assert "--include" in parser.format_help()

    args = parser.parse_args(["--label", "weekly", "--include", "Alpha"])
    script.__dict__["LABEL"] = args.label  # what the script's __main__ block does with it
    assert cli.main(script, args, transport=httpx.MockTransport(fake_services)) == 0
    assert [result.result for result in _reported(script)] == [{"label": "weekly"}]


@pytest.mark.integration
@pytest.mark.acceptance
def test_main_requires_a_tenant_selection(
    fake_msal: FakeMsal, caplog: pytest.LogCaptureFixture
) -> None:
    assert _main(LICENSED_USERS) == 2
    assert "Choose tenants with --include: a tenant, a group, or all." in caplog.text
    assert fake_msal.interactive_calls == 0


@pytest.mark.integration
def test_main_refuses_exclude_all_before_sign_in(
    fake_msal: FakeMsal, caplog: pytest.LogCaptureFixture
) -> None:
    assert _main(LICENSED_USERS, "--include", "Alpha", "--exclude", " ALL ") == 2
    assert "--exclude all would leave no tenant to run." in caplog.text
    assert fake_msal.interactive_calls == 0


# --- report() --------------------------------------------------------------------------


@pytest.mark.integration
@pytest.mark.acceptance
def test_report_gets_every_tenant_in_order_failures_included(
    fake_msal: FakeMsal, recorder: ModuleType, capsys: pytest.CaptureFixture[str]
) -> None:
    fake_msal.tenant_errors[TENANT_B] = CONSENT_REQUIRED
    assert _main(recorder, "--include", "all") == 1
    (alpha, beta) = _reported(recorder)
    assert (alpha.tenant.tenant_id, alpha.status, alpha.result) == (
        TENANT_A,
        "ok",
        [f"admin@{TENANT_A}.example"],
    )
    assert (beta.tenant.tenant_id, beta.status, beta.result) == (TENANT_B, "auth_failed", None)
    assert beta.error is not None
    assert "has no consent in tenant" in beta.error
    captured = capsys.readouterr()
    assert captured.out == ""  # the library writes nothing; stdout is the script's
    assert "Summary: ok=1, auth_failed=1, error=0, skipped=0" in captured.err


@pytest.mark.integration
def test_a_script_without_report_runs_and_writes_nothing(
    fake_msal: FakeMsal, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    script = _script(tmp_path, 'SCOPES = ["User.Read.All"]\n\ndef run(ctx):\n    return 1\n')
    assert _main(script, "--include", "all") == 0
    captured = capsys.readouterr()
    assert captured.out == ""
    assert "Summary: ok=2, auth_failed=0, error=0, skipped=0" in captured.err


@pytest.mark.integration
@pytest.mark.acceptance
def test_exit_2_when_report_fails(
    fake_msal: FakeMsal,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    script = _script(
        tmp_path,
        """
        SCOPES = ["User.Read.All"]

        def run(ctx):
            return 1

        def report(results):
            raise RuntimeError("disk full")
        """,
        name="audit",
    )
    assert _main(script, "--include", "all") == 2
    assert "audit's report() failed" in caplog.text
    assert "RuntimeError: disk full" in caplog.text
    assert "Summary: ok=2, auth_failed=0, error=0, skipped=0" in capsys.readouterr().err


@pytest.mark.integration
@pytest.mark.acceptance
def test_interrupt_skips_report(
    fake_msal: FakeMsal, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    script = _script(
        tmp_path,
        """
        SCOPES = ["User.Read.All"]
        REPORTED = []

        def run(ctx):
            if ctx.tenant.display_name == "beta corp":
                raise KeyboardInterrupt
            return 1

        def report(results):
            REPORTED.append(results)
        """,
    )
    assert _main(script, "--include", "all") == 130
    assert script.__dict__["REPORTED"] == []
    assert "Summary: ok=1, auth_failed=0, error=1, skipped=0" in capsys.readouterr().err


# --- warnings and errors again, after the summary --------------------------------------


@pytest.mark.integration
@pytest.mark.acceptance
def test_warnings_and_errors_are_printed_again_after_the_summary(
    fake_msal: FakeMsal,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO)
    script = _script(
        tmp_path,
        """
        import logging

        SCOPES = ["User.Read.All"]
        CONTINUE_ON_ERROR = True

        def run(ctx):
            logging.getLogger("script").info("checking")
            if ctx.tenant.display_name == "beta corp":
                raise RuntimeError("boom")
            logging.getLogger("script").warning("auditing is off")
        """,
    )
    assert _main(script, "--include", "all") == 1
    err = capsys.readouterr().err
    heading, *lines = err.split("Summary: ok=1, auth_failed=0, error=1, skipped=0\n")[
        1
    ].splitlines()
    assert heading == "Warnings and errors from this run (3):"
    # Each line as the terminal shows it, less the time.
    lines = [line.split(" ", 1)[1] for line in lines]
    path = tmp_path / "config" / "tenants.csv"
    assert lines[:2] == [
        f"WARNING [-] No tenants file at {path}; run m365-gdap-runner-customer-export to create it",
        "WARNING [Alpha] auditing is off",
    ]
    assert re.fullmatch(r"ERROR \[beta corp\] error in \d+\.\ds: RuntimeError\('boom'\)", lines[2])


@pytest.mark.integration
def test_a_run_that_logged_no_warning_prints_none_again(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, capsys: pytest.CaptureFixture[str]
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        ",".join(OWN_COLUMNS),
        f"{TENANT_A},Alpha,,Commercial,,TRUE,",
        f"{TENANT_B},beta corp,,Commercial,,TRUE,",
    )
    assert _main(recorder, "--include", "all", "--tenants-file", str(path)) == 0
    assert capsys.readouterr().err.endswith("Summary: ok=2, auth_failed=0, error=0, skipped=0\n")


@pytest.mark.integration
def test_a_run_that_stops_before_its_tenants_prints_no_warnings_again(
    fake_msal: FakeMsal, capsys: pytest.CaptureFixture[str]
) -> None:
    # The missing tenants file is a warning, and the unknown tenant ends the run before any
    # tenant runs: its error is the last line already.
    assert _main(LICENSED_USERS, "--include", "Contoso") == 2
    assert "Warnings and errors" not in capsys.readouterr().err


@pytest.mark.integration
def test_the_warning_log_is_removed_after_the_run(
    fake_msal: FakeMsal, recorder: ModuleType
) -> None:
    handlers = list(logging.getLogger().handlers)
    assert _main(recorder, "--include", "all") == 0
    assert logging.getLogger().handlers == handlers


# --- end to end ------------------------------------------------------------------------


@pytest.mark.integration
def test_runs_the_graph_example_and_its_report(
    fake_msal: FakeMsal, capsys: pytest.CaptureFixture[str]
) -> None:
    assert _main(LICENSED_USERS, "--include", "all", "--jobs", "2") == 0
    rows = json.loads(capsys.readouterr().out)
    assert [(r["tenantId"], r["tenant"], r["status"]) for r in rows] == [
        (TENANT_A, "Alpha", "ok"),
        (TENANT_B, "beta corp", "ok"),
    ]
    assert rows[1]["users"] == [
        {
            "userPrincipalName": f"admin@{TENANT_B}.example",
            "displayName": "Admin",
            "skuIds": "sku-1;sku-2",
        }
    ]
    assert fake_msal.interactive_calls == 1


@pytest.mark.integration
def test_runs_the_exchange_example_whose_log_is_the_report(
    fake_msal: FakeMsal, capsys: pytest.CaptureFixture[str], caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.INFO)
    assert _main(AUDIT_DISABLED, "--include", "all") == 0
    assert capsys.readouterr().out == ""
    # The runner's own warning, of the missing tenants file, aside.
    warnings = [
        r.getMessage()
        for r in caplog.records
        if r.levelno == logging.WARNING and r.name != cli.__name__
    ]
    assert warnings == ["Mailbox auditing is off tenant-wide (AuditDisabled is True)"]
    assert fake_msal.interactive_calls == 2


@pytest.mark.integration
def test_tenant_filters(fake_msal: FakeMsal, recorder: ModuleType) -> None:
    assert _main(recorder, "--include", "all", "--exclude", "BETA CORP") == 0
    assert _tenant_ids(recorder) == [TENANT_A]


@pytest.mark.integration
def test_exit_2_on_unknown_tenant_before_running_the_script(fake_msal: FakeMsal) -> None:
    seen: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request.url.path)
        return fake_services(request)

    assert _main(LICENSED_USERS, "--include", "Contoso", handler=handler) == 2
    assert "/v1.0/users" not in seen


@pytest.mark.integration
def test_exit_2_on_invalid_script_before_sign_in(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    script = _script(tmp_path, "def run(ctx):\n    return None\n")
    assert _main(script, "--include", "all") == 2
    assert "Script 'script' must define SCOPES" in caplog.text
    assert fake_msal.interactive_calls == 0


@pytest.mark.integration
def test_a_script_with_empty_scopes_gets_whatever_the_app_holds(
    fake_msal: FakeMsal, tmp_path: Path
) -> None:
    body = RECORDER.replace('SCOPES = ["User.Read.All"]', "SCOPES = []")
    script = _script(tmp_path, body)
    fake_msal.tenant_errors[TENANT_B] = CONSENT_REQUIRED
    assert _main(script, "--include", "all") == 1
    (alpha, beta) = _reported(script)
    assert (alpha.status, beta.status) == ("ok", "auth_failed")
    assert {scopes for _, scopes in fake_msal.silent_calls(TENANT_A)} == {
        ("https://graph.microsoft.com/.default",)
    }


def _grant_service(fake_msal: FakeMsal, patches: list[dict[str, str]]) -> Handler:
    """Fake services plus the app's one tenant-wide grant for Graph, in every tenant.

    The grant holds what the tenant has granted. A PATCH to it records the body and changes
    what the tenant has granted, as Entra then puts it in the tokens it issues.
    """

    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        tenant = bearer_tenant(request).lower()
        if path == "/v1.0/servicePrincipals":
            graph = GRAPH_APP_ID in request.url.params["$filter"]
            return httpx.Response(200, json={"value": [{"id": "graph" if graph else "app"}]})
        if path == "/v1.0/oauth2PermissionGrants":
            grant = {
                "id": "grant",
                "clientId": "app",
                "consentType": "AllPrincipals",
                "resourceId": "graph",
                "scope": fake_msal.granted[tenant],
            }
            return httpx.Response(200, json={"value": [grant]})
        if path == "/v1.0/oauth2PermissionGrants/grant" and request.method == "PATCH":
            body = json.loads(request.content)
            patches.append(body)
            fake_msal.granted[tenant] = body["scope"]
            return httpx.Response(204)
        return fake_services(request)

    return handler


@pytest.mark.integration
@pytest.mark.acceptance
def test_grant_when_possible_adds_the_missing_scopes_and_the_tenant_runs(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.INFO, logger="azure_auth")
    script = _script(tmp_path, RECORDER + "\n    GRANT_WHEN_POSSIBLE = True\n")
    fake_msal.granted[TENANT_A] = "DelegatedPermissionGrant.ReadWrite.All"
    patches: list[dict[str, str]] = []
    handler = _grant_service(fake_msal, patches)
    assert _main(script, "--include", "Alpha", handler=handler) == 0
    assert patches == [{"scope": "DelegatedPermissionGrant.ReadWrite.All User.Read.All"}]
    assert [result.status for result in _reported(script)] == ["ok"]
    assert "Added User.Read.All to the tenant-wide grant" in caplog.text


@pytest.mark.integration
def test_no_grant_when_possible_overrides_the_script(fake_msal: FakeMsal, tmp_path: Path) -> None:
    script = _script(tmp_path, RECORDER + "\n    GRANT_WHEN_POSSIBLE = True\n")
    fake_msal.granted[TENANT_A] = "DelegatedPermissionGrant.ReadWrite.All"
    patches: list[dict[str, str]] = []
    handler = _grant_service(fake_msal, patches)
    assert _main(script, "--include", "Alpha", "--no-grant-when-possible", handler=handler) == 1
    (alpha,) = _reported(script)
    assert alpha.status == "auth_failed"
    assert alpha.error is not None
    assert "could grant itself the requested scopes" in alpha.error
    assert patches == []


@pytest.mark.integration
def test_exit_2_on_missing_config(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    monkeypatch.delenv(ENV_PREFIX + "USERNAME")
    assert _main(LICENSED_USERS, "--include", "all") == 2
    assert "username" in caplog.text


@pytest.mark.integration
def test_exit_2_on_failed_sign_in(fake_msal: FakeMsal) -> None:
    fake_msal.interactive_result = {"error": "access_denied", "error_description": "cancelled"}
    assert _main(LICENSED_USERS, "--include", "all") == 2


# A script whose run() fails in Alpha, the first tenant, and works in beta corp.
FAILS_IN_ALPHA = """
    SCOPES = ["User.Read.All"]
    REPORTED = []

    def run(ctx):
        if ctx.tenant.display_name == "Alpha":
            raise RuntimeError("broken")
        return 1

    def report(results):
        REPORTED.append(results)
"""


@pytest.mark.integration
@pytest.mark.acceptance
@pytest.mark.parametrize(
    ("setting", "flags", "beta"),
    [
        ("", [], "skipped"),
        ("", ["--continue-on-error"], "ok"),
        ("CONTINUE_ON_ERROR = True", [], "ok"),
        ("CONTINUE_ON_ERROR = True", ["--no-continue-on-error"], "skipped"),
        ("CONTINUE_ON_ERROR = False", [], "skipped"),
    ],
    ids=["default", "flag", "script", "flag-over-script", "script-false"],
)
def test_an_error_in_run_stops_the_run_unless_the_script_or_a_flag_says_go_on(
    fake_msal: FakeMsal, tmp_path: Path, setting: str, flags: list[str], beta: str
) -> None:
    script = _script(tmp_path, FAILS_IN_ALPHA + "\n    " + setting + "\n")
    assert _main(script, "--include", "all", *flags) == 1
    assert [r.status for r in _reported(script)] == ["error", beta]


@pytest.mark.integration
@pytest.mark.acceptance
@pytest.mark.parametrize(
    ("env", "flags", "jobs"),
    [("", [], 1), ("3", [], 3), ("3", ["--jobs", "2"], 2)],
    ids=["default", "setting", "flag-over-setting"],
)
def test_jobs_come_from_the_setting_or_the_flag(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    recorder: ModuleType,
    env: str,
    flags: list[str],
    jobs: int,
) -> None:
    if env:
        monkeypatch.setenv(ENV_PREFIX + "JOBS", env)
    asked: dict[str, int] = {}
    check_access = PartnerSession.check_access

    def checking(self: PartnerSession, *args: Any, **kwargs: Any) -> Any:
        asked["checks"] = kwargs["jobs"]
        return check_access(self, *args, **kwargs)

    def running(*args: Any, **kwargs: Any) -> Any:
        asked["runs"] = kwargs["jobs"]
        return run_payload(*args, **kwargs)

    monkeypatch.setattr(PartnerSession, "check_access", checking)
    monkeypatch.setattr("m365_gdap_runner.cli.run_payload", running)
    assert _main(recorder, "--include", "all", *flags) == 0
    # The silent checks and the tenants' run() both work on that many at once.
    assert asked == {"checks": jobs, "runs": jobs}


@pytest.mark.integration
def test_exit_2_when_the_jobs_setting_is_below_1(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    monkeypatch.setenv(ENV_PREFIX + "JOBS", "0")
    assert _main(LICENSED_USERS, "--include", "all") == 2
    assert "jobs must be at least 1, not 0." in caplog.text
    assert fake_msal.interactive_calls == 0


@pytest.mark.integration
def test_undeclared_client_is_a_tenant_error(fake_msal: FakeMsal, tmp_path: Path) -> None:
    script = _script(
        tmp_path,
        """
        SCOPES = ["User.Read.All"]
        REPORTED = []

        def run(ctx):
            return ctx.exchange.run("Get-Mailbox")

        def report(results):
            REPORTED.append(results)
        """,
    )
    assert _main(script, "--include", "Alpha") == 1
    (result,) = _reported(script)
    assert result.status == "error"
    assert result.error is not None
    assert "ClientNotDeclared" in result.error


# --- the m365-gdap-runner-customer-export command --------------------------------------


@pytest.mark.integration
@pytest.mark.acceptance
def test_list_tenants_creates_the_tenants_file_and_prints_the_tenants(
    fake_msal: FakeMsal,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level(logging.INFO, logger=cli.__name__)
    assert _list_tenants() == 0
    assert capsys.readouterr().out.splitlines() == [f"{TENANT_A}\tAlpha", f"{TENANT_B}\tbeta corp"]
    path = tmp_path / "config" / "tenants.csv"
    assert path.read_text(encoding="utf-8-sig").splitlines() == [
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url",
        f"{TENANT_A},Alpha,,Commercial,,TRUE,",
        f"{TENANT_B},beta corp,,Commercial,,TRUE,",
    ]
    assert f"Created tenants file {path} with 2 tenant(s)" in caplog.text
    assert fake_msal.interactive_calls == 1


@pytest.mark.integration
def test_list_tenants_takes_the_settings_flags(fake_msal: FakeMsal, tmp_path: Path) -> None:
    path = tmp_path / "elsewhere" / "customers.csv"
    assert _list_tenants("--tenants-file", str(path)) == 0
    assert path.exists()
    assert not (tmp_path / "config" / "tenants.csv").exists()


@pytest.mark.integration
def test_list_tenants_exit_2_on_missing_config(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    monkeypatch.delenv(ENV_PREFIX + "USERNAME")
    assert _list_tenants() == 2
    assert "username" in caplog.text


@pytest.mark.integration
@pytest.mark.acceptance
def test_debug_log_gets_every_record_and_the_terminal_info_and_up(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    path = tmp_path / "logs" / "debug.log"
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_PATH", str(path))
    monkeypatch.setattr(cli, "configure_logging", CONFIGURE_LOGGING)
    for _run in range(2):  # each as its own process would, with handlers removed after
        with _restoring_loggers():
            assert _list_tenants() == 0
    log = path.read_text(encoding="utf-8")
    terminal = capsys.readouterr().err
    # This package's debug records, and a library's below WARNING, reach the file only.
    assert "DEBUG [-] m365_gdap_runner.session: Signing in to the partner tenant" in log
    assert "INFO [-] httpx: HTTP Request" in log
    assert "Signing in to the partner tenant" not in terminal
    assert "HTTP Request" not in terminal
    assert "Created tenants file" in terminal
    # The second run appends.
    assert log.count("Debug log started") == 2


@pytest.mark.integration
def test_a_debug_log_that_cannot_be_opened_is_a_warning_not_an_error(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    (tmp_path / "file").write_text("", encoding="utf-8")
    path = tmp_path / "file" / "debug.log"  # its folder is a file
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_PATH", str(path))
    assert _list_tenants() == 0
    assert f"Could not open debug log {path}" in caplog.text


@pytest.mark.integration
@pytest.mark.parametrize("trailing", ["", os.sep], ids=["existing", "new-with-separator"])
def test_a_debug_log_path_naming_a_folder_writes_the_default_file_name_in_it(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, trailing: str
) -> None:
    folder = tmp_path / "logs"
    if not trailing:
        folder.mkdir()
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_PATH", str(folder) + trailing)
    with _restoring_loggers():
        assert _list_tenants() == 0
    assert "Debug log started" in (folder / cli._DEBUG_LOG_NAME).read_text(encoding="utf-8")


@pytest.mark.unit
@pytest.mark.parametrize(
    ("setting", "expected"),
    [
        ("{tmp}", "{tmp}/m365_gdap_runner_debug.log"),
        ("{tmp}/new/", "{tmp}/new/m365_gdap_runner_debug.log"),
        ("{tmp}/new", "{tmp}/new"),
        ("{tmp}/debug.log", "{tmp}/debug.log"),
    ],
    ids=["existing-folder", "trailing-separator", "no-suffix-is-a-file", "file"],
)
def test_debug_log_file(tmp_path: Path, setting: str, expected: str) -> None:
    assert cli._debug_log_file(setting.format(tmp=tmp_path)) == Path(expected.format(tmp=tmp_path))


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_full_debug_log_is_compressed_and_only_the_newest_ones_kept(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    path = tmp_path / "debug.log"
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_PATH", str(path))
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_MAX_MB", "0.0005")  # about 500 bytes, a few lines
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_KEEP", "2")
    runs = 4  # each logs more than a file holds, so the oldest are deleted
    for _run in range(runs):
        with _restoring_loggers():
            assert _list_tenants() == 0
    assert sorted(p.name for p in tmp_path.glob("debug.log*")) == [
        "debug.log",
        "debug.log.1.gz",
        "debug.log.2.gz",
    ]
    assert path.stat().st_size <= 0.0005 * cli._MEGABYTE
    kept = path.read_text(encoding="utf-8")
    for number in (1, 2):
        old = gzip.decompress((tmp_path / f"debug.log.{number}.gz").read_bytes()).decode("utf-8")
        assert re.match(r"\d{4}-\d\d-\d\d \d\d:\d\d:\d\d,\d+ [A-Z]+ \[", old)
        kept += old
    assert kept.count("Debug log started") < runs


@pytest.mark.integration
def test_a_debug_log_with_no_size_limit_is_never_compressed(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    path = tmp_path / "debug.log"
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_PATH", str(path))
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_MAX_MB", "0")
    with _restoring_loggers():
        assert _list_tenants() == 0
    assert [p.name for p in tmp_path.glob("debug.log*")] == ["debug.log"]


@pytest.mark.integration
@pytest.mark.parametrize(
    ("name", "value", "message"),
    [
        ("DEBUG_LOG_MAX_MB", "-1", "debug_log_max_mb must not be negative, not -1.0."),
        ("DEBUG_LOG_KEEP", "0", "debug_log_keep must be at least 1, not 0."),
    ],
    ids=["negative-size", "keep-none"],
)
def test_a_bad_debug_log_limit_is_a_configuration_error(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    name: str,
    value: str,
    message: str,
) -> None:
    monkeypatch.setenv(ENV_PREFIX + "DEBUG_LOG_PATH", str(tmp_path / "debug.log"))
    monkeypatch.setenv(ENV_PREFIX + name, value)
    assert _list_tenants() == 2
    assert message in caplog.text
    assert not (tmp_path / "debug.log").exists()


@pytest.mark.unit
def test_a_debug_log_that_cannot_be_compressed_warns_once_and_keeps_growing(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    def held_open(self: cli._DebugLogHandler, source: str, dest: str) -> None:
        raise PermissionError(32, "The file is being used by another process", source)

    monkeypatch.setattr(cli._DebugLogHandler, "rotate", held_open)
    folder = tmp_path / "logs"
    folder.mkdir()
    path = folder / "debug.log"
    handler = cli._DebugLogHandler(path, max_bytes=200, keep=2)
    log = logging.getLogger("tests.debug_log")
    log.addHandler(handler)
    try:
        for number in range(20):
            log.warning("record %d", number)
    finally:
        log.removeHandler(handler)
        handler.close()
    assert caplog.text.count("Could not compress debug log") == 1
    assert [p.name for p in folder.iterdir()] == ["debug.log"]
    assert path.read_text(encoding="utf-8").splitlines() == [f"record {n}" for n in range(20)]


@pytest.mark.unit
def test_list_tenants_takes_no_selection_flags(capsys: pytest.CaptureFixture[str]) -> None:
    with pytest.raises(SystemExit) as info:
        cli.list_tenants(["--include", "all"])
    assert info.value.code == 2
    assert "unrecognized arguments: --include all" in capsys.readouterr().err


# --- tenants file ----------------------------------------------------------------------


def _tenants_csv(path: Path, *lines: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_run_warns_of_what_the_file_lacks_and_never_writes_it(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, caplog: pytest.LogCaptureFixture
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,pilot,notes",
        f"{TENANT_A},Alpha,alpha.example,Commercial,,TRUE,keep me",
        f"{TENANT_C},Gamma,,,,FALSE,",  # no GDAP relationship and a blank cloud
    )
    before = path.read_bytes()
    assert _main(recorder, "--include", "m365_gdap", "--tenants-file", str(path)) == 0
    # The file lacks m365_gdap, so it counts as TRUE for every GDAP tenant, beta corp included.
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]
    assert path.read_bytes() == before
    assert (
        f"Tenants file {path} lacks column(s) m365_gdap, m365_pw_url; run "
        "m365-gdap-runner-customer-export to add them" in caplog.text
    )
    assert (
        f"1 GDAP tenant(s) have no row in {path}; run m365-gdap-runner-customer-export to add "
        "them: beta corp" in caplog.text
    )
    assert (
        f"Looked up the blank m365_cloud of 1 tenant(s) in {path}; run "
        "m365-gdap-runner-customer-export to save them: Gamma" in caplog.text
    )


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_missing_tenants_file_is_a_warning_and_a_run_never_creates_it(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, caplog: pytest.LogCaptureFixture
) -> None:
    assert _main(recorder, "--include", "all") == 0
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]
    path = tmp_path / "config" / "tenants.csv"
    assert (
        f"No tenants file at {path}; run m365-gdap-runner-customer-export to create it"
        in caplog.text
    )
    assert not path.exists()


@pytest.mark.integration
@pytest.mark.acceptance
def test_list_tenants_appends_tenants_the_file_lacks_and_warns(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,pilot,notes",
        f"{TENANT_A},Alpha,alpha.example,Commercial,,TRUE,keep me",
    )
    assert _list_tenants("--tenants-file", str(path)) == 0
    # m365_gdap is new to the file, so it is TRUE for every GDAP tenant.
    assert path.read_text(encoding="utf-8-sig").splitlines() == [
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url,pilot,notes",
        f"{TENANT_A},Alpha,alpha.example,Commercial,,TRUE,,TRUE,keep me",
        f"{TENANT_B},beta corp,,Commercial,,TRUE,,,",
    ]
    assert f"Added 1 tenant(s) to {path}; they are in no group yet: beta corp" in caplog.text


@pytest.mark.integration
@pytest.mark.acceptance
def test_group_runs_only_the_marked_tenants(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,pilot",
        f"{TENANT_A},Alpha,,Commercial,FALSE",
        f"{TENANT_B},beta corp,,Commercial,TRUE",
        f"{TENANT_C},Gamma,,,FALSE",  # no GDAP relationship: harmless while unselected
    )
    assert _main(recorder, "--include", "pilot", "--tenants-file", str(path)) == 0
    assert _tenant_ids(recorder) == [TENANT_B]


@pytest.mark.integration
def test_groups_and_tenants_add_up_and_exclude_applies(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,pilot,patching",
        f"{TENANT_A},Alpha,,Commercial,TRUE,FALSE",
        f"{TENANT_B},beta corp,,Commercial,FALSE,TRUE",
    )
    file_flag = ["--tenants-file", str(path)]
    assert _main(recorder, "--include", "pilot", "--include", "PATCHING", *file_flag) == 0
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]
    assert _main(recorder, "--include", "pilot", "--include", "beta corp", *file_flag) == 0
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]
    assert _main(recorder, "--include", "pilot", "--exclude", "Alpha", *file_flag) == 0
    assert _tenant_ids(recorder) == []


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_tenant_may_have_more_than_one_row_while_they_agree(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, caplog: pytest.LogCaptureFixture
) -> None:
    header = f"{','.join(OWN_COLUMNS)},pilot"
    alpha = f"{TENANT_A},Alpha,,Commercial,,TRUE,"
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        header,
        f"{alpha},TRUE",
        f"{TENANT_B},beta corp,,Commercial,,TRUE,,FALSE",
        f"{alpha},FALSE",  # Alpha again: a group may differ
    )
    file_flag = ["--tenants-file", str(path)]
    assert _main(recorder, "--include", "pilot", *file_flag) == 0
    assert _tenant_ids(recorder) == [TENANT_A]
    _tenants_csv(
        path,
        header,
        f"{alpha},TRUE",
        f"{TENANT_A},Alpha,,Commercial,admin@alpha.example,TRUE,,TRUE",
    )
    assert _main(recorder, "--include", "pilot", *file_flag) == 2
    assert (
        f"rows 2 and 3 are both tenant {TENANT_A} but differ in column 'm365_username'"
        in caplog.text
    )


@pytest.mark.integration
def test_tenants_file_setting_locates_the_file(
    fake_msal: FakeMsal, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, recorder: ModuleType
) -> None:
    path = _tenants_csv(
        tmp_path / "elsewhere" / "customers.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,pilot",
        f"{TENANT_A},Alpha,,Commercial,TRUE",
    )
    monkeypatch.setenv(ENV_PREFIX + "TENANTS_FILE", str(path))
    assert _main(recorder, "--include", "pilot") == 0
    assert _tenant_ids(recorder) == [TENANT_A]
    assert not (tmp_path / "config" / "tenants.csv").exists()


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_workbook_tenants_file_picks_tenants_as_a_csv_does(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "tenants.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.append([*OWN_COLUMNS, "pilot"])
    sheet.append([TENANT_A, "Alpha", None, "Commercial", None, True, None, False])
    sheet.append([TENANT_B, "beta corp", None, "Commercial", None, True, None, True])
    sheet.add_table(Table(displayName="Tenants", ref="A1:H3"))
    workbook.save(path)
    before = path.read_bytes()
    assert _main(recorder, "--include", "pilot", "--tenants-file", str(path)) == 0
    assert _tenant_ids(recorder) == [TENANT_B]
    assert path.read_bytes() == before
    assert "WARNING" not in caplog.text


@pytest.mark.integration
@pytest.mark.acceptance
def test_list_tenants_creates_a_workbook_tenants_file(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    path = tmp_path / "tenants.xlsx"
    assert _list_tenants("--tenants-file", str(path)) == 0
    sheet = load_workbook(path)["Tenants"]
    assert [cell.value for cell in sheet["A"]] == ["m365_tenant_id", TENANT_A, TENANT_B]
    # A script then reads it as it reads a CSV.
    assert _main(recorder, "--include", "m365_gdap", "--tenants-file", str(path)) == 0
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]


def _gamma_csv(
    tmp_path: Path, cloud: str = "Commercial", username: str = GAMMA_ADMIN, gdap: str = ""
) -> Path:
    """A tenants file whose pilot group holds Alpha and Gamma, which has no GDAP relationship."""
    return _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,pilot",
        f"{TENANT_A},Alpha,,Commercial,,TRUE,TRUE",
        f"{TENANT_C},Gamma,,{cloud},{username},{gdap},TRUE",
    )


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_group_member_without_gdap_gets_a_sign_in_of_its_own(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.INFO, logger="m365_gdap_runner.session")
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        return fake_services(request)

    path = _gamma_csv(tmp_path)
    assert _main(recorder, "--include", "pilot", "--tenants-file", str(path), handler=handler) == 0
    assert [(r.tenant.tenant_id, r.status) for r in _reported(recorder)] == [
        (TENANT_A, "ok"),
        (TENANT_C, "ok"),
    ]
    # One browser sign-in for the partner, one for Gamma as its own account; none for Alpha.
    assert fake_msal.interactive_tenants() == [PARTNER_TENANT, TENANT_C]
    # Gamma's blank m365_gdap means the partner sign-in is never tried there.
    assert fake_msal.silent_calls(TENANT_C) == []
    assert f"Signing in to Gamma directly as {GAMMA_ADMIN}: its m365_gdap is not TRUE" in (
        caplog.text
    )
    users = [bearer_tenant(r) for r in seen if r.url.path == "/v1.0/users"]
    assert sorted(users) == [TENANT_A, TENANT_C]


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_tenant_in_another_cloud_is_signed_in_to_there(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    seen: list[httpx.Request] = []

    def handler(request: httpx.Request) -> httpx.Response:
        seen.append(request)
        return fake_services(request)

    path = _gamma_csv(tmp_path, cloud="USGov")
    assert _main(recorder, "--include", "Gamma", "--tenants-file", str(path), handler=handler) == 0
    (gamma,) = _reported(recorder)
    assert (gamma.tenant.tenant_id, gamma.tenant.cloud, gamma.status) == (TENANT_C, "USGov", "ok")
    # Straight to the US Government sign-in: the commercial partner sign-in was never tried.
    assert fake_msal.silent_calls(TENANT_C) == []
    assert fake_msal.authority(TENANT_C) == f"https://login.microsoftonline.us/{TENANT_C}"
    (users,) = [r for r in seen if r.url.path == "/v1.0/users"]
    assert users.url.host == "graph.microsoft.us"


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_gdap_tenant_missing_consent_gets_a_sign_in_of_its_own(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    fake_msal.tenant_errors[TENANT_B] = CONSENT_REQUIRED
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username",
        f"{TENANT_B},beta corp,,Commercial,admin@beta.example",
    )
    assert _main(recorder, "--include", "all", "--tenants-file", str(path)) == 0
    assert [r.status for r in _reported(recorder)] == ["ok", "ok"]
    # The browser opens at beta corp itself, as its own account, where consent is given.
    assert fake_msal.interactive_tenants() == [PARTNER_TENANT, TENANT_B]


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_gdap_tenant_marked_false_in_m365_gdap_skips_the_partner_sign_in(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    row = f"{TENANT_B},beta corp,,Commercial,admin@beta.example,FALSE,"
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url",
        row,
    )
    assert _main(recorder, "--include", "all", "--tenants-file", str(path)) == 0
    # beta corp is still a GDAP tenant, so all runs it, but only as its own account.
    assert [(r.tenant.tenant_id, r.status) for r in _reported(recorder)] == [
        (TENANT_A, "ok"),
        (TENANT_B, "ok"),
    ]
    assert fake_msal.silent_calls(TENANT_B) == []
    assert fake_msal.interactive_tenants() == [PARTNER_TENANT, TENANT_B]
    assert row in path.read_text(encoding="utf-8-sig").splitlines()


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_direct_sign_in_opens_the_password_page_in_the_tenants_file(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    opened: list[str] = []

    def open_page(url: str, *args: Any, **kwargs: Any) -> bool:
        opened.append(url)
        return True

    monkeypatch.setattr("webbrowser.open", open_page)
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url",
        f"{TENANT_A},Alpha,,Commercial,,TRUE,https://passwords.example/alpha",
        f"{TENANT_B},beta corp,,Commercial,admin@beta.example,FALSE,https://passwords.example/beta",
    )
    assert _main(recorder, "--include", "all", "--tenants-file", str(path)) == 0
    assert [r.status for r in _reported(recorder)] == ["ok", "ok"]
    # Only beta corp signs in directly; Alpha is reached through the partner sign-in.
    assert opened == ["https://passwords.example/beta"]


@pytest.mark.integration
@pytest.mark.acceptance
def test_a_tenant_with_no_gdap_and_no_username_is_left_out_with_a_warning(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, caplog: pytest.LogCaptureFixture
) -> None:
    # How a customer with no access at all, that Partner Center still lists, is marked.
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap",
        f"{TENANT_B},beta corp,,Commercial,,FALSE",
    )
    assert _main(recorder, "--include", "all", "--tenants-file", str(path)) == 0
    assert [(r.tenant.tenant_id, r.status) for r in _reported(recorder)] == [(TENANT_A, "ok")]
    assert (
        "Skipping 1 tenant(s) with neither m365_gdap TRUE nor an m365_username: beta corp"
        in caplog.text
    )
    assert [r for r in caplog.records if r.levelno >= logging.ERROR] == []
    assert fake_msal.silent_calls(TENANT_B) == []
    assert fake_msal.interactive_tenants() == [PARTNER_TENANT]


@pytest.mark.integration
@pytest.mark.regression
def test_a_blank_username_fails_the_tenant_without_a_sign_in_as_the_partner_account(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    fake_msal.tenant_errors[TENANT_C] = NOT_IN_TENANT  # the partner account is not in Gamma
    path = _gamma_csv(tmp_path, username="", gdap="TRUE")
    assert _main(recorder, "--include", "pilot", "--tenants-file", str(path)) == 1
    alpha, gamma = _reported(recorder)
    assert (alpha.status, gamma.status) == ("ok", "auth_failed")
    assert gamma.error is not None
    assert "AADSTS50020" in gamma.error
    assert "No direct sign-in to Gamma: its m365_username is blank." in gamma.error
    assert fake_msal.interactive_tenants() == [PARTNER_TENANT]


@pytest.mark.integration
@pytest.mark.acceptance
def test_include_m365_gdap_runs_only_the_tenants_it_marks(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    # Alpha is TRUE and Gamma blank; beta corp is appended, with TRUE.
    path = _gamma_csv(tmp_path)
    assert _main(recorder, "--include", "m365_gdap", "--tenants-file", str(path)) == 0
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]
    assert fake_msal.interactive_tenants() == [PARTNER_TENANT]


@pytest.mark.integration
def test_exit_2_when_a_name_is_both_a_tenant_and_a_group(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,Alpha",
        f"{TENANT_A},Alpha,,Commercial,TRUE",
    )
    assert _main(LICENSED_USERS, "--include", "alpha", "--tenants-file", str(path)) == 2
    assert f"'alpha' could mean tenant Alpha ({TENANT_A}) or group column 'Alpha'" in caplog.text


@pytest.mark.integration
def test_a_failed_direct_sign_in_fails_that_tenant_only(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    fake_msal.interactive_errors[TENANT_C] = {"error": "access_denied", "error_description": "no"}
    path = _gamma_csv(tmp_path)
    assert _main(recorder, "--include", "pilot", "--tenants-file", str(path)) == 1
    assert [r.status for r in _reported(recorder)] == ["ok", "auth_failed"]


@pytest.mark.integration
@pytest.mark.acceptance
def test_ctrl_c_then_enter_during_a_direct_sign_in_skips_that_tenant_only(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake_msal.interactive_interrupts.add(TENANT_C)
    monkeypatch.setattr("builtins.input", lambda: "")
    path = _gamma_csv(tmp_path)
    assert _main(recorder, "--include", "pilot", "--tenants-file", str(path)) == 1
    assert [(r.status, r.error) for r in _reported(recorder)] == [
        ("ok", None),
        ("auth_failed", "The sign-in to Gamma was skipped with Ctrl+C."),
    ]


@pytest.mark.integration
def test_all_runs_gdap_tenants_only(
    fake_msal: FakeMsal, tmp_path: Path, recorder: ModuleType
) -> None:
    path = _gamma_csv(tmp_path)
    assert _main(recorder, "--include", "all", "--tenants-file", str(path)) == 0
    assert _tenant_ids(recorder) == [TENANT_A, TENANT_B]


@pytest.mark.integration
def test_exit_2_when_the_groups_select_nothing(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,pilot",
        f"{TENANT_A},Alpha,,Commercial,FALSE",
    )
    assert _main(LICENSED_USERS, "--include", "pilot", "--tenants-file", str(path)) == 2
    assert "--include pilot selects no tenant." in caplog.text


@pytest.mark.integration
def test_list_tenants_fills_in_blank_clouds_by_discovery(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.INFO, logger=cli.__name__)
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,pilot",
        f"{TENANT_A},Alpha,,,,TRUE",
        f"{TENANT_B},beta corp,,,,FALSE",
        f"{TENANT_C},Gamma,,,,FALSE",  # no GDAP relationship, but a tenant all the same
    )
    assert _list_tenants("--tenants-file", str(path)) == 0
    assert path.read_text(encoding="utf-8-sig").splitlines() == [
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url,pilot",
        f"{TENANT_A},Alpha,,Commercial,,TRUE,,TRUE",
        f"{TENANT_B},beta corp,,Commercial,,TRUE,,FALSE",
        f"{TENANT_C},Gamma,,Commercial,,,,FALSE",
    ]
    assert f"Filled in the cloud of 3 tenant(s) in {path}: Alpha, beta corp, Gamma" in caplog.text


@pytest.mark.integration
@pytest.mark.acceptance
def test_list_tenants_adds_the_columns_the_file_lacks(
    fake_msal: FakeMsal, tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    caplog.set_level(logging.INFO, logger=cli.__name__)
    path = _tenants_csv(
        tmp_path / "tenants.csv",
        "m365_tenant_id,m365_tenant_name,pilot",
        f"{TENANT_A},Alpha,TRUE",
        f"{TENANT_B},beta corp,FALSE",
    )
    assert _list_tenants("--tenants-file", str(path)) == 0
    assert path.read_text(encoding="utf-8-sig").splitlines() == [
        "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,m365_username,m365_gdap,"
        "m365_pw_url,pilot",
        f"{TENANT_A},Alpha,,Commercial,,TRUE,,TRUE",
        f"{TENANT_B},beta corp,,Commercial,,TRUE,,FALSE",
    ]
    assert (
        "Added column(s) m365_domain, m365_cloud, m365_username, m365_gdap, m365_pw_url to "
        f"{path}" in caplog.text
    )


@pytest.mark.integration
@pytest.mark.parametrize(
    ("lines", "message"),
    [
        (
            ("m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud", f"{TENANT_A},Alpha,,"),
            "No tenant or group matches 'pilot'.",
        ),
        (
            (
                "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,pilot",
                f"{TENANT_A},Alpha,,,maybe",
            ),
            "row 2: 'maybe' in column 'pilot' is not TRUE, FALSE or blank",
        ),
        (
            (
                "m365_tenant_id,m365_tenant_name,m365_domain,m365_cloud,pilot",
                f"{TENANT_A},Alpha,,Mars,TRUE",
            ),
            "row 2: 'Mars' in column 'm365_cloud' is not a cloud",
        ),
    ],
)
def test_exit_2_on_a_bad_tenants_file(
    fake_msal: FakeMsal,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    lines: tuple[str, ...],
    message: str,
) -> None:
    path = _tenants_csv(tmp_path / "tenants.csv", *lines)
    assert _main(LICENSED_USERS, "--include", "pilot", "--tenants-file", str(path)) == 2
    assert message in caplog.text


@pytest.mark.integration
def test_unwritable_tenants_file_is_a_warning_not_an_error(
    fake_msal: FakeMsal,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    def refuse(self: TenantFile) -> None:
        raise PermissionError("held open by a spreadsheet")

    monkeypatch.setattr(TenantFile, "save", refuse)
    assert _list_tenants() == 0
    assert "Could not write tenants file" in caplog.text
    assert "held open by a spreadsheet" in caplog.text
    assert not (tmp_path / "config" / "tenants.csv").exists()
