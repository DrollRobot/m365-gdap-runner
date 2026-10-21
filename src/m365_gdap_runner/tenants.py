"""GDAP customer tenants, the tenants file, and how the command line picks tenants.

Graph says which customer tenants have a GDAP relationship (see ``session.py``). The tenants
file says which tenants belong to which group: a CSV, or an Excel workbook holding one table,
one row per customer, with the columns ``m365_tenant_id``, ``m365_tenant_name``,
``m365_domain``, ``m365_cloud``, ``m365_username``, ``m365_gdap`` and ``m365_pw_url``, and one
column per group holding TRUE or FALSE. ``--include NAME`` runs the tenants marked TRUE in
column ``NAME``, and ``--exclude NAME`` leaves them out (:func:`select`).

The file may be shared with tools that describe the same customers on other platforms, so
the columns this package owns carry the ``m365_`` prefix, every other column is left alone
unless a run names it as a group, and a row with no ``m365_tenant_id`` is a customer without
an M365 tenant: never selected, never an error.

Tenants are matched by ID only; the name and domain columns are for the person editing the
file. A column the file lacks is added, a GDAP tenant with no row is appended, and a blank
cloud is filled in by OpenID Connect discovery (:meth:`TenantFile.sync`). Only the
``m365-gdap-runner`` command saves those to the file; a script run fills them in memory and
warns. A row whose tenant has no GDAP relationship, or lives in another Microsoft cloud, can
still be selected: the runner signs in to it directly, as the account in ``m365_username``. A
blank cell means no direct sign-in. ``m365_pw_url``, when filled in, is the web page holding
that account's password, which opens with the direct sign-in's window.

``m365_gdap`` says which tenants the partner sign-in reaches, and is a group like any other,
so ``--include m365_gdap`` runs only those. TRUE tries the partner sign-in first. FALSE or
blank never tries it, and goes straight to the direct sign-in: GDAP can list a customer
whose tenant it gives no access to. With ``m365_username`` blank too, the tenant has no way
in, and a run leaves it out: that marks a customer with no access at all that cannot be
removed from Partner Center. An appended row gets TRUE, and so does every row GDAP lists
when the column itself is added.
"""

from __future__ import annotations

import contextlib
import csv
import os
import tempfile
import warnings
from collections.abc import Callable, Sequence
from dataclasses import dataclass, replace
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import BadZipFile

from azure_auth.clouds import CLOUDS, get_cloud
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from m365_gdap_runner.config import Settings
from m365_gdap_runner.config.paths import config_dir

# The columns this package owns. A new file gets them in this order, and an existing one gets
# any it lacks; other tools sharing the file add their own. ``m365_gdap``, and any column
# outside these, may serve as a group.
ID_COLUMN = "m365_tenant_id"
NAME_COLUMN = "m365_tenant_name"
DOMAIN_COLUMN = "m365_domain"
CLOUD_COLUMN = "m365_cloud"
USERNAME_COLUMN = "m365_username"
GDAP_COLUMN = "m365_gdap"
PW_URL_COLUMN = "m365_pw_url"
OWN_COLUMNS = (
    ID_COLUMN,
    NAME_COLUMN,
    DOMAIN_COLUMN,
    CLOUD_COLUMN,
    USERNAME_COLUMN,
    GDAP_COLUMN,
    PW_URL_COLUMN,
)

# Schemes an ``m365_pw_url`` may use. Anything else, such as a file path, is refused: opening
# it the way a web address is opened would run it.
PW_URL_SCHEMES = frozenset({"http", "https"})

# File name in the config directory when the ``tenants_file`` setting is empty.
TENANTS_FILE_NAME = "tenants.csv"

# The ``--include`` value that means every GDAP tenant, compared case-insensitively.
ALL = "all"

# Cell values that put a tenant in a group, or keep it out, compared case-insensitively:
# the spreadsheet booleans only. A blank cell, which every appended row has, keeps it out.
TRUE_VALUES = frozenset({"true"})
FALSE_VALUES = frozenset({"false", ""})


class TenantSelectionError(ValueError):
    """Raised when a selection names a tenant the runner cannot use, or selects nothing."""


class TenantFileError(ValueError):
    """Raised when the tenants file cannot be read or is not laid out as expected."""


@dataclass(frozen=True)
class Tenant:
    """A customer tenant: a GDAP customer, or a row of the tenants file.

    Attributes:
        tenant_id: Entra tenant ID.
        display_name: Tenant display name as set in Entra ID, or as the tenants file names it.
        cloud: The cloud it lives in, as azure-auth spells it. Filled in for a tenant a run
            selects; empty for one only listed.
        username: The account to sign in with when the partner sign-in cannot reach it; empty
            for no direct sign-in.
        gdap: False when the tenants file's ``m365_gdap`` cell is not TRUE, so the partner
            sign-in is never tried.
        pw_url: The web page holding *username*'s password, opened with the direct sign-in's
            window; empty for none.
    """

    tenant_id: str
    display_name: str
    cloud: str = ""
    username: str = ""
    gdap: bool = True
    pw_url: str = ""


@dataclass(frozen=True)
class TenantRow:
    """The M365 columns of one row of the tenants file, stripped of surrounding whitespace.

    Attributes:
        tenant_id: The ``m365_tenant_id`` cell.
        display_name: The ``m365_tenant_name`` cell; may be empty.
        cloud: The ``m365_cloud`` cell as azure-auth spells the cloud's name, or empty.
        username: The ``m365_username`` cell, or empty when blank or the column is absent.
        gdap: True when the ``m365_gdap`` cell is TRUE; False when it is FALSE, blank, or the
            column is absent.
        pw_url: The ``m365_pw_url`` cell, or empty when blank or the column is absent.
    """

    tenant_id: str
    display_name: str
    cloud: str
    username: str = ""
    gdap: bool = False
    pw_url: str = ""


@dataclass(frozen=True)
class SyncReport:
    """What :meth:`TenantFile.sync` changed.

    Attributes:
        columns_added: Columns this package owns that the file lacked, in the order added.
        added: Tenants that got a new row.
        clouds_filled: Tenants whose existing row got its ``m365_cloud`` filled in.
    """

    columns_added: list[str]
    added: list[Tenant]
    clouds_filled: list[Tenant]

    @property
    def changed(self) -> bool:
        """Whether the file differs from what was loaded."""
        return bool(self.columns_added or self.added or self.clouds_filled)


def tenants_file_path(settings: Settings) -> Path:
    """Return where the tenants file lives.

    Args:
        settings: Resolved settings.

    Returns:
        The ``tenants_file`` setting with ``~`` expanded, or ``tenants.csv`` in the config
        directory when the setting is empty.
    """
    if settings.tenants_file:
        return Path(settings.tenants_file).expanduser()
    return config_dir() / TENANTS_FILE_NAME


def _is_workbook(path: Path) -> bool:
    """Return True if *path* names an Excel workbook rather than a CSV.

    Args:
        path: The tenants file.

    Returns:
        Whether its name ends in ``.xlsx``, in any case.
    """
    return path.suffix.lower() == ".xlsx"


def _cell_text(value: object) -> str:
    """Return a workbook cell's value as the text a CSV cell would hold.

    Args:
        value: The cell's value as openpyxl reads it; a formula's is its last calculated one.

    Returns:
        TRUE or FALSE for Excel's booleans, ``""`` for an empty cell, else the value as text.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _cell_value(text: str) -> str | bool | None:
    """Return what to write to a workbook cell for a cell's text, undoing :func:`_cell_text`.

    Args:
        text: The cell's text.

    Returns:
        Excel's boolean for TRUE or FALSE, None (an empty cell) for ``""``, else the text.
    """
    if text in ("TRUE", "FALSE"):
        return text == "TRUE"
    return text or None


def _open_workbook(path: Path, *, data_only: bool) -> tuple[Workbook, Worksheet, Table]:
    """Open an Excel workbook and find its one table, on whichever sheet it is.

    Args:
        path: The ``.xlsx`` file.
        data_only: True reads each formula's last calculated value; False keeps the formulas,
            to write the workbook back.

    Returns:
        The workbook, the sheet holding the table, and the table.

    Raises:
        TenantFileError: If the workbook cannot be read, or holds no table or more than one.
    """
    try:
        # openpyxl warns of workbook features it does not support; they are not this
        # package's to report.
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            workbook = load_workbook(path, data_only=data_only)
    except (OSError, BadZipFile, InvalidFileException, KeyError) as exc:
        raise TenantFileError(f"Cannot read {path}: {exc}") from exc
    tables = [
        (sheet, table)
        for sheet in workbook.worksheets
        if isinstance(sheet, Worksheet)
        for table in sheet.tables.values()
    ]
    if len(tables) != 1:
        raise TenantFileError(
            f"{path} has {len(tables)} tables; a tenants workbook must have exactly one."
        )
    sheet, table = tables[0]
    return workbook, sheet, table


class TenantFile:
    """The tenants file: one row per customer, one TRUE/FALSE column per group.

    Rows keep every cell as the text the file holds, a workbook's as :func:`_cell_text` gives
    it, so :meth:`save` writes back what was read plus any added columns and rows. Only the
    group column a caller asks about, and the cloud of the rows marked in it, are checked;
    every other cell may hold anything.

    Attributes:
        path: Where the file is, or will be, on disk.
        columns: Column names in file order.
        rows: One mapping per row, column name to cell text.
        first_row: The spreadsheet row number of the first of *rows*, for error messages.
    """

    def __init__(
        self,
        path: Path,
        columns: Sequence[str],
        rows: Sequence[dict[str, str]] = (),
        *,
        first_row: int = 2,
    ) -> None:
        """Hold a file's contents in memory.

        Args:
            path: Where the file is, or will be, on disk.
            columns: Column names in file order. :meth:`sync` adds any of
                :data:`OWN_COLUMNS` missing.
            rows: One mapping per row, column name to cell text.
            first_row: The spreadsheet row number of the first of *rows*: 2, below a header
                in row 1, unless a workbook's table starts lower.
        """
        self.path = path
        self.columns = list(columns)
        self.rows = [dict(row) for row in rows]
        self.first_row = first_row
        # As read, so saving a workbook writes only the cells that changed.
        self._read_columns = list(columns)
        self._read_rows = [dict(row) for row in rows]

    @classmethod
    def empty(cls, path: Path) -> TenantFile:
        """Return a file with this package's columns and no rows, not yet on disk.

        Args:
            path: Where :meth:`save` will write it.

        Returns:
            The empty file.
        """
        return cls(path, OWN_COLUMNS)

    @classmethod
    def load(cls, path: Path) -> TenantFile:
        """Read the file.

        A CSV with nothing in it, not even a header, counts as :meth:`empty`. A header that
        lacks some of this package's columns is fine; :meth:`sync` adds them.

        Args:
            path: The file to read: an Excel workbook when its name ends in ``.xlsx`` (see
                :meth:`_load_workbook`), else a CSV in UTF-8, with or without a byte-order
                mark.

        Returns:
            The file's contents.

        Raises:
            TenantFileError: If the file cannot be read, a CSV is not UTF-8, or a workbook
                does not hold exactly one table.
        """
        if _is_workbook(path):
            return cls._load_workbook(path)
        try:
            with path.open(newline="", encoding="utf-8-sig") as handle:
                reader = csv.DictReader(handle, restval="")
                columns = reader.fieldnames
                # A row longer than the header puts the extra cells under the key None.
                rows = [
                    {key: value for key, value in row.items() if key is not None} for row in reader
                ]
        except UnicodeDecodeError as exc:
            raise TenantFileError(f"{path} is not UTF-8 text: {exc}") from exc
        except OSError as exc:
            raise TenantFileError(f"Cannot read {path}: {exc}") from exc
        if columns is None:
            return cls.empty(path)
        return cls(path, columns, rows)

    @classmethod
    def _load_workbook(cls, path: Path) -> TenantFile:
        """Read the one table of an Excel workbook, on whichever sheet it is.

        The table's header row names the columns. A totals row is left out.

        Args:
            path: The ``.xlsx`` file.

        Returns:
            The table's contents, each cell as :func:`_cell_text` gives it.

        Raises:
            TenantFileError: If the workbook cannot be read, or holds no table or more than
                one.
        """
        _, sheet, table = _open_workbook(path, data_only=True)
        cells = CellRange(table.ref)
        header, *body = sheet.iter_rows(
            min_row=cells.min_row,
            max_row=cells.max_row - (table.totalsRowCount or 0),
            min_col=cells.min_col,
            max_col=cells.max_col,
            values_only=True,
        )
        columns = [_cell_text(value) for value in header]
        rows = [dict(zip(columns, map(_cell_text, row), strict=True)) for row in body]
        return cls(path, columns, rows, first_row=cells.min_row + 1)

    @property
    def groups(self) -> list[str]:
        """Columns that look like groups, in file order.

        A column that may be a group (see :meth:`_group_candidates`) whose filled cells are
        all TRUE or FALSE. Columns other tools keep in a shared file are left out. This is the
        hint when a name matches nothing; :meth:`find_group` itself takes any candidate.
        """
        return [column for column in self._group_candidates() if self._is_boolean(column)]

    def _group_candidates(self) -> list[str]:
        """Return every named column that may be a group, in file order.

        Returns:
            ``m365_gdap``, and every column this package does not own.
        """
        return [
            column
            for column in self.columns
            if column and (column == GDAP_COLUMN or column not in OWN_COLUMNS)
        ]

    def _is_boolean(self, column: str) -> bool:
        """Return True if every cell of *column* is TRUE, FALSE or blank.

        Args:
            column: A column name.

        Returns:
            Whether the column could be a group.
        """
        allowed = TRUE_VALUES | FALSE_VALUES
        return all(row.get(column, "").strip().casefold() in allowed for row in self.rows)

    def find_group(self, name: str) -> str | None:
        """Find the group column a name means, case-insensitively.

        Args:
            name: The name given on the command line.

        Returns:
            The column name as the file spells it, or None when no column matches.

        Raises:
            TenantFileError: If more than one column matches.
        """
        needle = name.strip().casefold()
        matches = [column for column in self._group_candidates() if column.casefold() == needle]
        if len(matches) > 1:
            raise TenantFileError(
                f"{name!r} matches {len(matches)} columns of {self.path}: {', '.join(matches)}."
            )
        return matches[0] if matches else None

    def _group_column(self, group: str) -> str:
        """Resolve a group name to its column, case-insensitively.

        Args:
            group: The name given on the command line.

        Returns:
            The column name as the file spells it.

        Raises:
            TenantFileError: If no column, or more than one, matches.
        """
        column = self.find_group(group)
        if column is None:
            known = ", ".join(self.groups) or "none"
            raise TenantFileError(f"{self.path} has no group column {group!r}. Groups: {known}.")
        return column

    def members(self, group: str) -> list[TenantRow]:
        """Return the rows marked TRUE in a group column, in file order.

        A marked row with no ``m365_tenant_id`` is a customer without an M365 tenant and is
        left out.

        Args:
            group: The group's column name, in any case.

        Returns:
            The member rows.

        Raises:
            TenantFileError: If there is no such column, a cell in it is not TRUE, FALSE or
                blank, or a marked row's ``m365_cloud`` names no known cloud, its
                ``m365_gdap`` is not TRUE, FALSE or blank, or its ``m365_pw_url`` is not a web
                address.
        """
        column = self._group_column(group)
        chosen: list[TenantRow] = []
        for number, row in enumerate(self.rows, start=self.first_row):
            cell = row.get(column, "")
            value = cell.strip().casefold()
            if value in FALSE_VALUES:
                continue
            if value not in TRUE_VALUES:
                raise TenantFileError(
                    f"{self.path} row {number}: {cell!r} in column {column!r} is not TRUE, "
                    "FALSE or blank."
                )
            member = self._row(row, number)
            if member.tenant_id:
                chosen.append(member)
        return chosen

    def _row(self, row: dict[str, str], number: int) -> TenantRow:
        """Pick the M365 columns out of a row.

        Args:
            row: Column name to cell text.
            number: The row's number in the spreadsheet, for error messages.

        Returns:
            The row's tenant ID, display name, cloud, username, ``m365_gdap`` and
            ``m365_pw_url``, stripped; the cloud spelled as azure-auth names it.

        Raises:
            TenantFileError: If ``m365_cloud`` is filled but names no known cloud,
                ``m365_gdap`` is not TRUE, FALSE or blank, or ``m365_pw_url`` is filled but is
                not an http or https address.
        """
        cell = row.get(CLOUD_COLUMN, "").strip()
        cloud = ""
        if cell:
            try:
                cloud = get_cloud(cell).name
            except ValueError:
                names = ", ".join(known.name for known in CLOUDS)
                raise TenantFileError(
                    f"{self.path} row {number}: {cell!r} in column {CLOUD_COLUMN!r} is not a "
                    f"cloud. Use {names}, or blank."
                ) from None
        gdap = row.get(GDAP_COLUMN, "")
        if gdap.strip().casefold() not in TRUE_VALUES | FALSE_VALUES:
            raise TenantFileError(
                f"{self.path} row {number}: {gdap!r} in column {GDAP_COLUMN!r} is not TRUE, "
                "FALSE or blank."
            )
        pw_url = row.get(PW_URL_COLUMN, "").strip()
        address = urlsplit(pw_url)
        if pw_url and (address.scheme.lower() not in PW_URL_SCHEMES or not address.netloc):
            raise TenantFileError(
                f"{self.path} row {number}: {pw_url!r} in column {PW_URL_COLUMN!r} is not a web "
                "address. Use an http or https address, or blank."
            )
        return TenantRow(
            row.get(ID_COLUMN, "").strip(),
            row.get(NAME_COLUMN, "").strip(),
            cloud,
            row.get(USERNAME_COLUMN, "").strip(),
            gdap=gdap.strip().casefold() in TRUE_VALUES,
            pw_url=pw_url,
        )

    def extra_tenants(self, gdap: Sequence[Tenant]) -> list[Tenant]:
        """Return the tenants the file has a row for but GDAP does not list, in file order.

        Args:
            gdap: The GDAP tenants.

        Returns:
            One tenant per ID, named as the file names it, or by its ID when unnamed.
        """
        seen = {tenant.tenant_id.lower() for tenant in gdap}
        extra: list[Tenant] = []
        for row in self.rows:
            tenant_id = row.get(ID_COLUMN, "").strip()
            if tenant_id and tenant_id.lower() not in seen:
                seen.add(tenant_id.lower())
                extra.append(Tenant(tenant_id, row.get(NAME_COLUMN, "").strip() or tenant_id))
        return extra

    def details(self, tenant: Tenant, *, default_cloud: str) -> Tenant:
        """Return *tenant* with the cloud, sign-in details and ``m365_gdap`` its row gives.

        Args:
            tenant: A selected tenant.
            default_cloud: The cloud when the row has none, or there is no row: the partner's.

        Returns:
            The tenant, with ``cloud`` always filled in.

        Raises:
            TenantFileError: If the row's ``m365_cloud`` names no known cloud, its
                ``m365_gdap`` is not TRUE, FALSE or blank, or its ``m365_pw_url`` is not a web
                address.
        """
        key = tenant.tenant_id.lower()
        for number, row in enumerate(self.rows, start=self.first_row):
            if row.get(ID_COLUMN, "").strip().lower() == key:
                found = self._row(row, number)
                return replace(
                    tenant,
                    cloud=found.cloud or default_cloud,
                    username=found.username,
                    gdap=found.gdap,
                    pw_url=found.pw_url,
                )
        return replace(tenant, cloud=default_cloud)

    def _add_missing_columns(self) -> list[str]:
        """Add the columns of :data:`OWN_COLUMNS` the file lacks, blank in every row.

        They go right after the last column this package owns, so its columns stay together,
        or at the end when the file has none of them.

        Returns:
            The columns added, in :data:`OWN_COLUMNS` order.
        """
        missing = [column for column in OWN_COLUMNS if column not in self.columns]
        if not missing:
            return []
        own = [index for index, column in enumerate(self.columns) if column in OWN_COLUMNS]
        at = own[-1] + 1 if own else len(self.columns)
        self.columns[at:at] = missing
        for row in self.rows:
            for column in missing:
                row.setdefault(column, "")
        return missing

    def sync(self, tenants: Sequence[Tenant], cloud_of: Callable[[str], str]) -> SyncReport:
        """Add missing columns, append a row for each tenant the file lacks, and fill in clouds.

        A column of :data:`OWN_COLUMNS` the file lacks is added, blank in every row, except
        that ``m365_gdap`` is TRUE in the rows of the GDAP tenants. An appended row gets the
        tenant's ID and name, its cloud from *cloud_of*, TRUE in ``m365_gdap``, and every
        other column blank. Any other row with a tenant ID and a blank ``m365_cloud``, a GDAP
        tenant or not, gets it from *cloud_of* too, so a file another tool wrote, or that
        predates the column, fills in over time. Nothing is written; call :meth:`save`.

        Args:
            tenants: The GDAP tenants.
            cloud_of: Returns the name of the cloud a tenant ID lives in, or ``""`` when it
                cannot be found. Called once per tenant that needs it.

        Returns:
            What changed.
        """
        gdap = {tenant.tenant_id.lower(): tenant for tenant in tenants}
        columns_added = self._add_missing_columns()
        if GDAP_COLUMN in columns_added:
            for existing in self.rows:
                if existing.get(ID_COLUMN, "").strip().lower() in gdap:
                    existing[GDAP_COLUMN] = "TRUE"
        by_id: dict[str, dict[str, str]] = {}
        for existing in self.rows:
            by_id.setdefault(existing.get(ID_COLUMN, "").strip().lower(), existing)
        known = dict(by_id)
        added: list[Tenant] = []
        for tenant in tenants:
            key = tenant.tenant_id.lower()
            if key not in by_id:
                row = dict.fromkeys(self.columns, "")
                row[ID_COLUMN] = tenant.tenant_id
                row[NAME_COLUMN] = tenant.display_name
                row[CLOUD_COLUMN] = cloud_of(tenant.tenant_id)
                row[GDAP_COLUMN] = "TRUE"
                self.rows.append(row)
                by_id[key] = row
                added.append(tenant)
        filled: list[Tenant] = []
        for key, row in known.items():
            if not key or row.get(CLOUD_COLUMN, "").strip():
                continue
            tenant_id = row[ID_COLUMN].strip()
            cloud = cloud_of(tenant_id)
            if cloud:
                row[CLOUD_COLUMN] = cloud
                named = row.get(NAME_COLUMN, "").strip() or tenant_id
                filled.append(gdap.get(key) or Tenant(tenant_id, named))
        return SyncReport(columns_added, added, filled)

    def save(self) -> None:
        """Write the file, replacing it in one step so a failure leaves the old one intact.

        The parent directory is created if needed. A CSV is written whole, with a UTF-8
        byte-order mark so spreadsheets open it correctly. A workbook gets only what changed
        (see :meth:`_updated_workbook`).

        Raises:
            OSError: If the file or its directory cannot be written, for example while a
                spreadsheet holds the file open.
            TenantFileError: If a workbook cannot be read, or its table cannot grow.
        """
        workbook = self._updated_workbook() if _is_workbook(self.path) else None
        self.path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temp = tempfile.mkstemp(
            dir=self.path.parent, prefix=self.path.name + ".", suffix=".tmp"
        )
        try:
            if workbook is None:
                with os.fdopen(descriptor, "w", newline="", encoding="utf-8-sig") as handle:
                    writer = csv.DictWriter(handle, fieldnames=self.columns, extrasaction="ignore")
                    writer.writeheader()
                    writer.writerows(self.rows)
            else:
                os.close(descriptor)
                workbook.save(temp)
            os.replace(temp, self.path)
        except BaseException:
            with contextlib.suppress(OSError):
                os.unlink(temp)
            raise

    def _updated_workbook(self) -> Workbook:
        """Return the workbook with this file's columns and rows in its table, for :meth:`save`.

        A new workbook gets one sheet and one table, both named ``Tenants``. In an existing
        one, only the cells that differ from what :meth:`load` read are written: a new column
        goes right of the table, a new row below it, and the table grows to hold them. Every
        other cell keeps its value, formula and formatting. TRUE and FALSE are written as
        Excel's booleans.

        Returns:
            The workbook, not yet saved.

        Raises:
            TenantFileError: If the workbook cannot be read or holds no table or more than
                one, or its table must grow but has a totals row or cells in the way.
        """
        if not self.path.exists():
            return self._new_workbook()
        workbook, sheet, table = _open_workbook(self.path, data_only=False)
        cells = CellRange(table.ref)
        top, left = cells.min_row, cells.min_col
        new_columns = [column for column in self.columns if column not in self._read_columns]
        order = [*self._read_columns, *new_columns]
        last_row = top + len(self.rows)
        last_column = left + len(order) - 1
        if len(self.rows) > len(self._read_rows) or new_columns:
            if table.totalsRowCount:
                raise TenantFileError(
                    f"Cannot add to the table in {self.path}: it has a totals row. Turn the "
                    "totals row off, then run m365-gdap-runner again."
                )
            below = sheet.iter_rows(
                min_row=cells.max_row + 1,
                max_row=last_row,
                min_col=left,
                max_col=last_column,
                values_only=True,
            )
            beside = sheet.iter_rows(
                min_row=top,
                max_row=cells.max_row,
                min_col=cells.max_col + 1,
                max_col=last_column,
                values_only=True,
            )
            if any(value is not None for row in [*below, *beside] for value in row):
                raise TenantFileError(
                    f"Cannot add to the table in {self.path}: the cells it would grow into, "
                    "below it or to its right, are not empty."
                )
            ref = f"{get_column_letter(left)}{top}:{get_column_letter(last_column)}{last_row}"
            table.ref = ref
            if table.autoFilter is not None:
                table.autoFilter.ref = ref
            next_id = max((column.id for column in table.tableColumns), default=0) + 1
            for offset, column in enumerate(new_columns):
                sheet.cell(top, left + order.index(column), column)
                table.tableColumns.append(TableColumn(id=next_id + offset, name=column))
        at = {column: left + index for index, column in enumerate(order)}
        for number, row in enumerate(self.rows):
            before = self._read_rows[number] if number < len(self._read_rows) else {}
            for column, column_at in at.items():
                value = row.get(column, "")
                if value != before.get(column, ""):
                    sheet.cell(top + 1 + number, column_at, _cell_value(value))
        return workbook

    def _new_workbook(self) -> Workbook:
        """Return a new workbook holding this file's columns and rows in one table.

        Returns:
            A workbook with one sheet and one table, both named ``Tenants``.
        """
        workbook = Workbook()
        for sheet in list(workbook.worksheets):
            workbook.remove(sheet)
        sheet = workbook.create_sheet("Tenants")
        sheet.append(self.columns)
        for row in self.rows:
            sheet.append([_cell_value(row.get(column, "")) for column in self.columns])
        # A table needs a row under its header, even a blank one.
        last_row = max(len(self.rows), 1) + 1
        table = Table(
            displayName="Tenants",
            ref=f"A1:{get_column_letter(len(self.columns))}{last_row}",
            # Excel's default for a new table.
            tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
        )
        sheet.add_table(table)
        return workbook


def _named(
    spec: str, gdap: Sequence[Tenant], known: Sequence[Tenant], tenant_file: TenantFile
) -> set[str]:
    """Resolve one ``--include`` or ``--exclude`` value to the tenants it names.

    A tenant ID always names its tenant. Any other value must mean exactly one thing:
    ``all``, one tenant's display name, or one group column.

    Args:
        spec: ``all``, a tenant's ID or display name, or a group column's name; any case.
        gdap: The GDAP tenants, which ``all`` means.
        known: Every tenant a value may name: the GDAP tenants and the file's other rows.
        tenant_file: The tenants file, for groups.

    Returns:
        The lower-cased IDs of the tenants named.

    Raises:
        TenantFileError: If two columns match the value, or it names a group whose column
            holds a cell that is not TRUE, FALSE or blank, or a member row is not valid.
        TenantSelectionError: If the value names nothing, or could mean more than one thing.
    """
    needle = spec.strip().casefold()
    for tenant in known:
        if tenant.tenant_id.casefold() == needle:
            return {tenant.tenant_id.lower()}
    by_name = [t for t in known if t.display_name.casefold() == needle]
    column = tenant_file.find_group(spec)
    meanings = [f"tenant {t.display_name} ({t.tenant_id})" for t in by_name]
    if needle == ALL:
        meanings.insert(0, "every GDAP tenant")
    if column is not None:
        meanings.append(f"group column {column!r} of {tenant_file.path}")
    if not meanings:
        groups = ", ".join(tenant_file.groups) or "none"
        raise TenantSelectionError(
            f"No tenant or group matches {spec!r}. Groups in {tenant_file.path}: {groups}."
        )
    if len(meanings) > 1:
        fixes = ["name the tenant by its ID"] if by_name else []
        if column is not None:
            fixes.append("rename the column")
        raise TenantSelectionError(
            f"{spec!r} could mean {' or '.join(meanings)}; {' or '.join(fixes)}."
        )
    if needle == ALL:
        return {tenant.tenant_id.lower() for tenant in gdap}
    if column is not None:
        return {row.tenant_id.lower() for row in tenant_file.members(column)}
    return {by_name[0].tenant_id.lower()}


def select(
    gdap: Sequence[Tenant],
    tenant_file: TenantFile,
    include: Sequence[str],
    exclude: Sequence[str] = (),
) -> list[Tenant]:
    """Pick the tenants to run: every tenant *include* names, less every one *exclude* names.

    A value is ``all`` for every GDAP tenant, a tenant's ID or display name (a GDAP customer
    or a row of the tenants file), or a group column, for the rows marked TRUE in it; any
    case. An exclusion beats an inclusion. A value that names nothing is an error, never a
    silent no-op, and so is one that could mean two things: a name a tenant and a group
    share, ``all`` as a tenant's or group's name, or two tenants' shared name.

    Args:
        gdap: The GDAP tenants.
        tenant_file: The tenants file.
        include: The values naming the tenants to run.
        exclude: The values naming the tenants to leave out.

    Returns:
        The selected tenants: the GDAP tenants first, in their order, then the file's other
        rows, in file order.

    Raises:
        TenantFileError: If two columns match a value, or a group named holds a cell that is
            not TRUE, FALSE or blank, or a member row is not valid.
        TenantSelectionError: If a value names nothing or could mean two things, or
            *include* selects no tenant.
    """
    known = [*gdap, *tenant_file.extra_tenants(gdap)]
    included: set[str] = set()
    for spec in include:
        included |= _named(spec, gdap, known, tenant_file)
    if not included:
        raise TenantSelectionError(f"--include {', '.join(include)} selects no tenant.")
    for spec in exclude:
        included -= _named(spec, gdap, known, tenant_file)
    return [tenant for tenant in known if tenant.tenant_id.lower() in included]
